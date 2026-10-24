"""Purchase register and expense register - the auditor's Excel books for a
period (Reports, the admin console and the monthly pack).

  Purchase register: one row per purchase bill per GST rate, with the
  supplier's invoice no./date and GSTIN, taxable value, tax and the input
  tax credit taken; goods returned to suppliers (debit notes) as minus
  rows. Margin-scheme and blocked lines show their tax with no credit.
  Expense register: every posted expense voucher, with its GST bill
  details and credit."""

from collections import defaultdict
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import selectinload

from app.models.expense import PAYMENT_MODE_LABELS, Expense, ExpenseStatus
from app.models.purchase import ITC_LABELS, ItcEligibility, PurchaseBill, PurchaseReturn, PurchaseStatus

ZERO = Decimal("0.00")
HEADS = ("taxable_value", "cgst_amount", "sgst_amount", "igst_amount")


def _rate_rows(lines, eligible_credit: bool):
    """Per-rate sums of a document's lines, and the credit in each."""
    rates = defaultdict(lambda: {**{h: ZERO for h in HEADS}, "itc": ZERO})
    for line in lines:
        r = rates[Decimal(line.gst_rate)]
        for h in HEADS:
            r[h] += Decimal(getattr(line, h))
        if eligible_credit and line.itc_eligibility == ItcEligibility.ELIGIBLE:
            r["itc"] += Decimal(line.cgst_amount) + Decimal(line.sgst_amount) + Decimal(line.igst_amount)
    return sorted(rates.items())


def purchase_register_rows(tenant, start: date, end: date, gstin_id: int | None = None) -> list[dict]:
    rows = []
    bills = PurchaseBill.query.filter(PurchaseBill.tenant_id == tenant.id, PurchaseBill.status == PurchaseStatus.POSTED,
                                      PurchaseBill.bill_date >= start, PurchaseBill.bill_date <= end)
    returns = PurchaseReturn.query.filter(PurchaseReturn.tenant_id == tenant.id,
                                          PurchaseReturn.status == PurchaseStatus.POSTED,
                                          PurchaseReturn.return_date >= start, PurchaseReturn.return_date <= end)
    if gstin_id:
        bills, returns = bills.filter(PurchaseBill.gstin_id == gstin_id), returns.filter(PurchaseReturn.gstin_id == gstin_id)
    for bill in bills.options(selectinload(PurchaseBill.lines)).order_by(PurchaseBill.bill_date, PurchaseBill.number):
        snap = bill.supplier_snapshot or {}
        for i, (rate, r) in enumerate(_rate_rows(bill.lines, bill.itc_cgst + bill.itc_sgst + bill.itc_igst > 0)):
            rows.append({"date": bill.bill_date, "number": bill.number, "kind": "Purchase bill",
                         "supplier_invoice": bill.supplier_invoice_no, "supplier_invoice_date": bill.supplier_invoice_date,
                         "supplier": snap.get("name", ""), "gstin": snap.get("gstin") or "", "state": snap.get("state_code") or "",
                         "rcm": bill.is_rcm, "rate": rate, **r, "total": Decimal(bill.grand_total) if i == 0 else None,
                         "sign": 1})
    for ret in returns.options(selectinload(PurchaseReturn.lines), selectinload(PurchaseReturn.bill)).order_by(
            PurchaseReturn.return_date, PurchaseReturn.number):
        snap = ret.supplier_snapshot or {}
        for i, (rate, r) in enumerate(_rate_rows(ret.lines, ret.itc_cgst + ret.itc_sgst + ret.itc_igst > 0)):
            rows.append({"date": ret.return_date, "number": ret.number, "kind": "Debit note (return)",
                         "supplier_invoice": ret.bill.supplier_invoice_no if ret.bill else "",
                         "supplier_invoice_date": ret.bill.supplier_invoice_date if ret.bill else None,
                         "supplier": snap.get("name", ""), "gstin": snap.get("gstin") or "",
                         "state": snap.get("state_code") or "", "rcm": False, "rate": rate,
                         **{k: -v for k, v in r.items()}, "total": -Decimal(ret.grand_total) if i == 0 else None,
                         "sign": -1})
    rows.sort(key=lambda r: (r["date"], r["number"]))
    return rows


def _sheet(ws, title_lines, headers, rows):
    for t in title_lines:
        ws.append([t])
    ws.append([])
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column[len(title_lines) + 1:] if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 40)


def _f(v):
    return float(v) if v is not None else None


def purchase_register_workbook(tenant, start: date, end: date, gstin_id: int | None = None) -> BytesIO:
    rows = purchase_register_rows(tenant, start, end, gstin_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase register"
    body = [[r["date"].isoformat(), r["number"], r["kind"], r["supplier_invoice"],
             r["supplier_invoice_date"].isoformat() if r["supplier_invoice_date"] else "", r["supplier"], r["gstin"],
             r["state"], "Yes" if r["rcm"] else "", _f(r["rate"]), _f(r["taxable_value"]), _f(r["cgst_amount"]),
             _f(r["sgst_amount"]), _f(r["igst_amount"]), _f(r["itc"]), _f(r["total"])] for r in rows]
    totals = {k: sum((r[k] for r in rows), ZERO) for k in (*HEADS, "itc")}
    body.append(["Total", "", "", "", "", "", "", "", "", None, _f(totals["taxable_value"]), _f(totals["cgst_amount"]),
                 _f(totals["sgst_amount"]), _f(totals["igst_amount"]), _f(totals["itc"]),
                 _f(sum((r["total"] for r in rows if r["total"] is not None), ZERO))])
    _sheet(ws, [f"Purchase register - {tenant.display_name}", f"{start.isoformat()} to {end.isoformat()}"],
           ["Date", "Our no.", "Type", "Supplier invoice", "Invoice date", "Supplier", "GSTIN", "State", "Reverse charge",
            "GST rate %", "Taxable value", "CGST", "SGST", "IGST", "ITC taken", "Bill total"], body)
    ws.append([])
    ws.append(["ITC is the tax on eligible lines only - blocked (s.17(5)), ineligible and margin-scheme lines carry none."])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def expense_register_rows(tenant, start: date, end: date, gstin_id: int | None = None) -> list[Expense]:
    q = Expense.query.filter(Expense.tenant_id == tenant.id, Expense.status == ExpenseStatus.POSTED,
                             Expense.expense_date >= start, Expense.expense_date <= end)
    if gstin_id:
        q = q.filter(Expense.gstin_id == gstin_id)
    return q.order_by(Expense.expense_date, Expense.number).all()


def expense_register_workbook(tenant, start: date, end: date, gstin_id: int | None = None) -> BytesIO:
    rows = expense_register_rows(tenant, start, end, gstin_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Expense register"
    body = [[e.expense_date.isoformat(), e.number, e.category.name, e.payee, e.supplier_gstin or "", e.bill_number or "",
             "Yes" if e.is_rcm else "", _f(e.gst_rate) if e.is_gst_bill else None,
             _f(e.taxable_value) if e.is_gst_bill else None, _f(e.cgst_amount), _f(e.sgst_amount), _f(e.igst_amount),
             _f(e.itc_total), ITC_LABELS[e.itc_eligibility] if e.is_gst_bill else "",
             PAYMENT_MODE_LABELS.get(e.payment_mode, ""), _f(e.amount)] for e in rows]
    body.append(["Total", "", "", "", "", "", "", None, _f(sum((e.taxable_value for e in rows if e.is_gst_bill), ZERO)),
                 _f(sum((e.cgst_amount for e in rows), ZERO)), _f(sum((e.sgst_amount for e in rows), ZERO)),
                 _f(sum((e.igst_amount for e in rows), ZERO)), _f(sum((e.itc_total for e in rows), ZERO)), "", "",
                 _f(sum((e.amount for e in rows), ZERO))])
    _sheet(ws, [f"Expense register - {tenant.display_name}", f"{start.isoformat()} to {end.isoformat()}"],
           ["Date", "Voucher", "Category", "Paid to", "Supplier GSTIN", "Bill no.", "Reverse charge", "GST rate %",
            "Taxable value", "CGST", "SGST", "IGST", "ITC taken", "Credit", "Paid by", "Amount"], body)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
