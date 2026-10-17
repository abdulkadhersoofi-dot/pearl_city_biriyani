"""Excel/PDF/JSON serialization for the GSTR-1/3B working papers computed
in app.reports.gstr. Kept separate from that module so the aggregation
logic (the part that has to be *correct*) stays readable without the
formatting/openpyxl/WeasyPrint plumbing mixed in.
"""

import json
from datetime import date
from decimal import Decimal
from io import BytesIO

from flask import current_app, render_template
from openpyxl import Workbook
from openpyxl.styles import Font
from weasyprint import HTML

HEADER_FONT = Font(bold=True)


def _autosize(ws):
    for col in ws.columns:
        length = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(length + 2, 10), 40)


def _write_header(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT


def _money(v) -> float:
    return float(v or 0)


def build_gstr1_workbook(data: dict) -> BytesIO:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary.append([f"GSTR-1 working paper - {data['gstin'].display_label}"])
    summary.append([f"Period: {data['period']} ({data['period_start']} to {data['period_end']})"])
    summary.append([])
    _write_header(summary, 4, ["Taxable value", "CGST", "SGST", "IGST", "Total"])
    t = data["totals"]
    summary.append([_money(t["taxable_value"]), _money(t["cgst"]), _money(t["sgst"]), _money(t["igst"]), _money(t["total"])])
    _autosize(summary)

    b2b = wb.create_sheet("B2B")
    _write_header(b2b, 1, ["Customer GSTIN", "Customer name", "Document type", "Document no.", "Date", "Place of supply", "Taxable value", "CGST", "SGST", "IGST"])
    row = 2
    for entry in data["b2b"]:
        for d in entry["documents"]:
            b2b.append([
                entry["customer_gstin"], entry["customer_name"], d["doc_type"], d["doc_number"],
                d["doc_date"].isoformat(), d["place_of_supply"],
                _money(d["taxable_value"]), _money(d["cgst"]), _money(d["sgst"]), _money(d["igst"]),
            ])
            row += 1
    _autosize(b2b)

    b2c = wb.create_sheet("B2CS")
    _write_header(b2c, 1, ["Place of supply", "GST rate %", "Taxable value", "CGST", "SGST", "IGST"])
    for b in data["b2c"]:
        b2c.append([b.place_of_supply, _money(b.gst_rate), _money(b.taxable_value), _money(b.cgst), _money(b.sgst), _money(b.igst)])
    _autosize(b2c)

    cdn = wb.create_sheet("Credit-Debit Notes")
    _write_header(cdn, 1, ["Note no.", "Date", "Type", "Against invoice", "Customer GSTIN", "Customer name", "Taxable value", "CGST", "SGST", "IGST", "Total"])
    for n in data["credit_debit_notes"]:
        cdn.append([
            n["note_number"], n["note_date"].isoformat(), n["document_type"], n["against_invoice"],
            n["customer_gstin"] or "", n["customer_name"],
            _money(n["taxable_value"]), _money(n["cgst"]), _money(n["sgst"]), _money(n["igst"]), _money(n["total"]),
        ])
    _autosize(cdn)

    b2cl = wb.create_sheet("B2CL")
    _write_header(b2cl, 1, ["Invoice no.", "Date", "Customer", "Place of supply", "GST rate %", "Taxable value", "IGST", "Invoice value"])
    for d in data.get("b2cl", []):
        for r in d["rates"]:
            b2cl.append([d["doc_number"], d["doc_date"].isoformat(), d["customer_name"], d["place_of_supply"],
                         _money(r["gst_rate"]), _money(r["taxable_value"]), _money(r["igst"]), _money(d["total"])])
    _autosize(b2cl)

    for title, key in (("HSN B2B", "hsn_b2b"), ("HSN B2C", "hsn_b2c")):
        sheet = wb.create_sheet(title)
        _write_header(sheet, 1, ["HSN/SAC", "GST rate %", "Unit", "Qty", "Taxable value", "CGST", "SGST", "IGST", "Total", "Note"])
        for h in data.get(key, []):
            sheet.append([h.hsn, _money(h.gst_rate), h.unit, float(h.qty), _money(h.taxable_value), _money(h.cgst),
                          _money(h.sgst), _money(h.igst), _money(h.total),
                          "Margin scheme - value other than margin (Rule 32(5))" if h.margin_remainder else ""])
        _autosize(sheet)

    hsn = wb.create_sheet("HSN Summary")
    _write_header(hsn, 1, ["HSN/SAC", "GST rate %", "Unit", "Qty", "Taxable value", "CGST", "SGST", "IGST", "Total", "Note"])
    for h in data["hsn_summary"]:
        hsn.append([h.hsn, _money(h.gst_rate), h.unit, float(h.qty), _money(h.taxable_value), _money(h.cgst), _money(h.sgst), _money(h.igst), _money(h.total),
                    "Margin scheme - value other than margin (Rule 32(5))" if h.margin_remainder else ""])
    _autosize(hsn)

    if "advances" in data:
        adv = wb.create_sheet("Advances 11A-11B")
        _write_header(adv, 1, ["Table", "Place of supply", "GST rate %", "Taxable value", "CGST", "SGST", "IGST"])
        for table, rows in (("11A received", data["advances"]["received"]), ("11B adjusted / refunded", data["advances"]["adjusted"])):
            for b in rows:
                adv.append([table, b.place_of_supply, _money(b.gst_rate), _money(b.taxable_value), _money(b.cgst),
                            _money(b.sgst), _money(b.igst)])
        _autosize(adv)

    if "nil_exempt" in data:
        nil = wb.create_sheet("Nil-Exempt (8)")
        _write_header(nil, 1, ["Description", "Nil rated", "Exempted", "Non-GST"])
        for row in data["nil_exempt"]["rows"].values():
            nil.append([row["label"], _money(row["nil_rated"]), _money(row["exempted"]), _money(row["non_gst"])])
        _autosize(nil)

    docs = wb.create_sheet("Documents issued")
    _write_header(docs, 1, ["Document type", "Series prefix", "From", "To", "Total issued", "Cancelled"])
    for s in data["document_summary"]:
        docs.append([s.label, s.prefix, s.from_number, s.to_number, s.total_count, s.cancelled_count])
    _autosize(docs)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_gstr3b_workbook(data: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "GSTR-3B"
    ws.append([f"GSTR-3B working paper - {data['gstin'].display_label}"])
    ws.append([f"Period: {data['period']} ({data['period_start']} to {data['period_end']})"])
    ws.append([])
    _write_header(ws, 4, ["3.1 Nature of supply", "Taxable value", "CGST", "SGST", "IGST"])
    rows = [
        ("(a) Outward taxable supplies (other than zero-rated, nil-rated, exempted)", data["outward_taxable"]),
        ("(b) Outward taxable supplies (zero-rated)", data["outward_zero_rated"]),
        ("(c) Other outward supplies (nil-rated, exempted)", data["outward_nil_rated"]),
    ]
    for label, b in rows:
        ws.append([label, _money(b["taxable_value"]), _money(b["cgst"]), _money(b["sgst"]), _money(b["igst"])])
    ws.append([])
    ws.append(["Gross output tax payable (before ITC)", _money(data["gross_tax_payable"])])
    if "advances" in data:
        net = data["advances"]["net"]
        ws.append(["3.1(a) includes tax on advances received less advances adjusted / refunded (GSTR-1 11A - 11B)",
                   _money(net["taxable_value"]), _money(net["cgst"]), _money(net["sgst"]), _money(net["igst"])])
    ws.append([])
    itc = data.get("itc")
    if itc is not None:
        _write_header(ws, ws.max_row + 1, ["3.1(d) Inward supplies liable to reverse charge", "Taxable value", "CGST", "SGST", "IGST"])
        r = itc["rcm_inward"]
        ws.append(["Reverse-charge purchases", _money(r["taxable_value"]), _money(r["cgst"]), _money(r["sgst"]), _money(r["igst"])])
        ws.append([])
        _write_header(ws, ws.max_row + 1, ["4. Eligible ITC", "", "CGST", "SGST", "IGST"])
        for label, key in (("(A)(3) Inward supplies liable to reverse charge", "a3_rcm"), ("(A)(5) All other ITC", "a5_other"),
                           ("(B)(1) Reversed - ineligible / blocked u/s 17(5)", "b1_ineligible"),
                           ("(B)(2) Reversed - goods returned to suppliers", "b2_returns"), ("(C) Net ITC available", "net")):
            row = itc[key]
            ws.append([label, "", _money(row["cgst"]), _money(row["sgst"]), _money(row["igst"])])
        ws.append([])
        pay = data.get("payment")
        if pay is not None:
            ws.append([])
            _write_header(ws, ws.max_row + 1, ["6.1 Tax payable", "Output tax", "Paid by IGST credit", "Paid by CGST credit",
                                               "Paid by SGST credit", "Output tax in cash", "Reverse charge (cash)", "Total cash"])
            for head in ("igst", "cgst", "sgst"):
                ws.append([head.upper(), _money(pay["output"][head]), _money(pay["utilised"]["igst"][head]),
                           _money(pay["utilised"]["cgst"][head]), _money(pay["utilised"]["sgst"][head]),
                           _money(pay["output_in_cash"][head]), _money(pay["rcm_in_cash"][head]), _money(pay["cash"][head])])
            ws.append(["Net payable in cash", "", "", "", "", "", "", _money(pay["cash_total"])])
            ws.append(["Credit carried forward", _money(pay["carry_forward"]["igst"]), _money(pay["carry_forward"]["cgst"]),
                       _money(pay["carry_forward"]["sgst"])])
        ws.append(["Claim only the ITC your GSTR-2B shows. Tax on reverse-charge purchases is payable in cash."])
    else:
        ws.append(["Input Tax Credit is not tracked in this system - this app records sales only."])
        ws.append(["Add eligible ITC from purchase records before computing net cash payable."])
    ws.append(["Reverse Charge Invoices are excluded above - the recipient pays that tax directly, not this tenant."])
    _autosize(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _json_default(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, date):
        return value.isoformat()
    if hasattr(value, "__dict__"):
        return {k: v for k, v in vars(value).items() if not k.startswith("_")}
    return str(value)


def _strip_non_serializable(data: dict) -> dict:
    out = dict(data)
    out.pop("tenant", None)
    gstin = out.pop("gstin", None)
    if gstin is not None:
        out["gstin"] = {"gstin": gstin.gstin, "state_code": gstin.state_code, "state_name": gstin.state_name}
    return out


def to_json_bytes(data: dict) -> bytes:
    return json.dumps(_strip_non_serializable(data), default=_json_default, indent=2).encode("utf-8")


def build_sales_workbook(data: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append([f"Sales report - {data['tenant'].display_name}"])
    ws.append([f"{data['start'].isoformat()} to {data['end'].isoformat()}"])
    ws.append([])
    _write_header(
        ws,
        4,
        [
            "Date", "Document type", "Document no.", "Billing GSTIN", "Customer", "Customer GSTIN",
            "Payment mode", "Taxable value", "CGST", "SGST", "IGST", "Total",
        ],
    )
    for r in data["rows"]:
        ws.append([
            r.doc_date.isoformat(), r.doc_type, r.doc_number, r.billing_gstin, r.customer_name,
            r.customer_gstin or "", (r.payment_mode or "").upper(),
            _money(r.sign * r.taxable_value), _money(r.sign * r.cgst), _money(r.sign * r.sgst),
            _money(r.sign * r.igst), _money(r.signed_total),
        ])
    ws.append([])
    t = data["totals"]
    ws.append(["", "", "", "", "", "", "Total", _money(t["taxable_value"]), _money(t["cgst"]), _money(t["sgst"]), _money(t["igst"]), _money(t["total"])])
    _autosize(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def render_sales_pdf(data: dict) -> bytes:
    html = render_template("reports/sales_pdf.html", data=data, firm_name=data["tenant"].display_name)
    return HTML(string=html, base_url=current_app.root_path).write_pdf()


def render_gstr1_pdf(data: dict) -> bytes:
    html = render_template("reports/gstr1_pdf.html", data=data, firm_name=current_app.config["FIRM_NAME"])
    return HTML(string=html, base_url=current_app.root_path).write_pdf()


def render_gstr3b_pdf(data: dict) -> bytes:
    html = render_template("reports/gstr3b_pdf.html", data=data, firm_name=current_app.config["FIRM_NAME"])
    return HTML(string=html, base_url=current_app.root_path).write_pdf()
