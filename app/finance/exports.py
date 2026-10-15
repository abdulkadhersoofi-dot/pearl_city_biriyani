"""Excel and PDF for the P&L, item gross profit and daily summary."""

from io import BytesIO

from flask import current_app, render_template
from openpyxl import Workbook
from openpyxl.styles import Font
from weasyprint import HTML


def _bold_last(ws):
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)


def _widths(ws, first=3):
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column[first:] if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 12), 48)


def pnl_workbook(st: dict, location_name: str) -> BytesIO:
    cur, prior = st["current"], st["prior"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Profit & loss"
    ws.append([f"Profit & loss - {st['tenant'].display_name} - {location_name}"])
    ws.append([f"{cur['start'].isoformat()} to {cur['end'].isoformat()}"
               + (" (cash-basis estimate - no stock records)" if cur["basis"] == "cash" else "")])
    ws.append([])
    head = ["", "This period"] + ([f"{st['prior_span'][0].isoformat()} to {st['prior_span'][1].isoformat()}"] if prior else [])
    ws.append(head)
    _bold_last(ws)

    def row(label, amount, prior_amount=None, bold=False):
        ws.append([label, float(amount)] + ([float(prior_amount)] if prior is not None and prior_amount is not None else []))
        if bold:
            _bold_last(ws)

    for sec in st["sections"]:
        ws.append([sec["title"]])
        _bold_last(ws)
        for line in sec["lines"]:
            row("   " + line["label"], line["amount"], line["prior"])
        row(f"Total {sec['title'].lower()}", sec["total"], sec["prior_total"], bold=True)
        if sec["key"] == "cost":
            row("Gross profit", cur["gross_profit"], prior["gross_profit"] if prior else None, bold=True)
            ws.append(["Gross profit %", float(cur["gp_pct"]) if cur["gp_pct"] is not None else None]
                      + ([float(prior["gp_pct"]) if prior["gp_pct"] is not None else None] if prior else []))
    row("Net profit", cur["net_profit"], prior["net_profit"] if prior else None, bold=True)
    _widths(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def render_pnl_pdf(st: dict, location_name: str) -> bytes:
    html = render_template("reports/profit_loss_pdf.html", st=st, location_name=location_name,
                           firm_name=st["tenant"].display_name)
    return HTML(string=html, base_url=current_app.root_path).write_pdf()


def item_profit_workbook(data: dict, tenant) -> BytesIO:
    wb = Workbook()
    for title, rows in (("By item", data["rows"]), ("By category", data["categories"])):
        ws = wb.active if title == "By item" else wb.create_sheet(title)
        ws.title = title
        ws.append([f"Gross profit {title.lower()} - {tenant.display_name}"])
        ws.append([f"{data['start'].isoformat()} to {data['end'].isoformat()}"])
        ws.append([])
        ws.append(["Item" if title == "By item" else "Category", "Category" if title == "By item" else "", "Qty",
                   "Sales (excl. GST)", "Cost", "Gross profit", "GP %"])
        _bold_last(ws)
        for r in rows:
            ws.append([r.name, r.category if title == "By item" else "", float(r.qty), float(r.sales), float(r.cost),
                       float(r.gross_profit), float(r.gp_pct) if r.gp_pct is not None else None])
        t = data["totals"]
        ws.append(["Total", "", None, float(t.sales), float(t.cost), float(t.gross_profit),
                   float(t.gp_pct) if t.gp_pct is not None else None])
        _bold_last(ws)
        _widths(ws)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def render_daily_pdf(data: dict, tenant) -> bytes:
    html = render_template("reports/daily_summary_pdf.html", data=data, firm_name=tenant.display_name)
    return HTML(string=html, base_url=current_app.root_path).write_pdf()
