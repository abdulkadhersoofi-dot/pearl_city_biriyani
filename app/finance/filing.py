"""Filing tracker and the monthly pack (Phase 7).

Due dates: GSTR-1 on gst_config.GSTR1_DUE_DAY and GSTR-3B on
GSTR3B_DUE_DAY of the month after the period (monthly filers).

Monthly pack: everything the firm needs to file and audit one client's
month, in one ZIP - GSTR-1 (Excel, PDF, JSON), GSTR-3B (Excel, PDF,
JSON), the HSN summary, the purchase and expense registers (when the
client records them), the P&L and the Tally import files."""

import calendar
from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font

from app.extensions import db
from app.models.filing import RETURN_LABELS, FilingStatus, ReturnType
from app.models.invoice import Invoice
from app.models.mixins import utcnow
from app.models.pos_bill import POSBill
from app.utils.gst_config import GSTR1_DUE_DAY, GSTR3B_DUE_DAY

DUE_DAYS = {ReturnType.GSTR1: GSTR1_DUE_DAY, ReturnType.GSTR3B: GSTR3B_DUE_DAY}


def period_bounds(period: str) -> tuple[date, date]:
    year, month = (int(x) for x in period.split("-"))
    return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])


def due_date(period: str, return_type: ReturnType) -> date:
    _, end = period_bounds(period)
    nxt = date(end.year + (end.month == 12), end.month % 12 + 1, 1)
    return nxt.replace(day=DUE_DAYS[return_type])


def status_row(tenant_id: int, gstin_id: int, period: str, return_type: ReturnType, create: bool = False):
    row = FilingStatus.query.filter_by(tenant_id=tenant_id, gstin_id=gstin_id, period=period,
                                       return_type=return_type).first()
    if row is None and create:
        row = FilingStatus(tenant_id=tenant_id, gstin_id=gstin_id, period=period, return_type=return_type)
        db.session.add(row)
    return row


def document_count(tenant_id: int, gstin_id: int, period: str) -> int:
    start, end = period_bounds(period)
    from app.finance.pnl import day_start
    from datetime import timedelta

    invoices = Invoice.query.filter(Invoice.tenant_id == tenant_id, Invoice.gstin_id == gstin_id,
                                    Invoice.invoice_date >= start, Invoice.invoice_date <= end).count()
    bills = POSBill.query.filter(POSBill.tenant_id == tenant_id, POSBill.gstin_id == gstin_id,
                                 POSBill.completed_at >= day_start(start),
                                 POSBill.completed_at < day_start(end + timedelta(days=1))).count()
    return invoices + bills


def tracker_rows(tenants, period: str, today: date | None = None) -> list[dict]:
    today = today or date.today()
    rows = []
    for tenant in tenants:
        for gstin in tenant.gstins.filter_by(is_active=True).all():
            docs = document_count(tenant.id, gstin.id, period)
            for rt in ReturnType:
                st = status_row(tenant.id, gstin.id, period, rt)
                due = due_date(period, rt)
                rows.append({"tenant": tenant, "gstin": gstin, "return_type": rt, "label": RETURN_LABELS[rt],
                             "due": due, "documents": docs, "status": st,
                             "overdue": not (st and st.filed) and today > due})
    return rows


def mark_ready(tenant_id, gstin_id, period, user, ready: bool = True) -> list[FilingStatus]:
    out = []
    for rt in ReturnType:
        row = status_row(tenant_id, gstin_id, period, rt, create=True)
        if not row.filed:
            row.data_ready = ready
            row.data_ready_at = utcnow() if ready else None
            row.data_ready_by_id = user.id if ready else None
        out.append(row)
    db.session.flush()
    return out


def mark_filed(tenant_id, gstin_id, period, return_type: ReturnType, user, filed: bool, on: date | None, arn: str | None):
    row = status_row(tenant_id, gstin_id, period, return_type, create=True)
    row.filed = filed
    row.filed_on = (on or date.today()) if filed else None
    row.arn = ((arn or "").strip()[:30] or None) if filed else None
    row.filed_by_id = user.id if filed else None
    db.session.flush()
    return row


# --- the monthly pack ---------------------------------------------------------------

def _hsn_workbook(g1) -> BytesIO:
    wb = Workbook()
    for i, (title, key) in enumerate((("HSN all", "hsn_summary"), ("HSN B2B", "hsn_b2b"), ("HSN B2C", "hsn_b2c"))):
        ws = wb.active if i == 0 else wb.create_sheet(title)
        ws.title = title
        ws.append(["HSN/SAC", "GST rate %", "Unit", "Qty", "Taxable value", "CGST", "SGST", "IGST", "Total"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for h in g1.get(key, []):
            ws.append([h.hsn, float(h.gst_rate), h.unit, float(h.qty), float(h.taxable_value), float(h.cgst),
                       float(h.sgst), float(h.igst), float(h.total)])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def pack_files(tenant, period: str) -> list[tuple[str, bytes]]:
    from app.finance.exports import pnl_workbook
    from app.finance.pnl import statement
    from app.finance.tally import tally_zip
    from app.models.tenant import RegistrationType
    from app.reports.export import (build_gstr1_workbook, build_gstr3b_workbook, render_gstr1_pdf,
                                    render_gstr3b_pdf, to_json_bytes)
    from app.reports.gstr import gstr1_data, gstr3b_data
    from app.reports.registers import expense_register_workbook, purchase_register_workbook

    start, end = period_bounds(period)
    files = []
    if tenant.registration_type == RegistrationType.REGULAR:
        for gstin in tenant.gstins.filter_by(is_active=True).all():
            tag = gstin.gstin or f"gstin-{gstin.id}"
            g1, g3 = gstr1_data(tenant, gstin, period), gstr3b_data(tenant, gstin, period)
            files += [(f"{tag}/GSTR1_{period}.xlsx", build_gstr1_workbook(g1).getvalue()),
                      (f"{tag}/GSTR1_{period}.pdf", render_gstr1_pdf(g1)),
                      (f"{tag}/GSTR1_{period}.json", to_json_bytes(g1)),
                      (f"{tag}/GSTR3B_{period}.xlsx", build_gstr3b_workbook(g3).getvalue()),
                      (f"{tag}/GSTR3B_{period}.pdf", render_gstr3b_pdf(g3)),
                      (f"{tag}/GSTR3B_{period}.json", to_json_bytes(g3)),
                      (f"{tag}/HSN_summary_{period}.xlsx", _hsn_workbook(g1).getvalue())]
    if tenant.has_feature("purchases"):
        files.append((f"Purchase_register_{period}.xlsx", purchase_register_workbook(tenant, start, end).getvalue()))
    if tenant.has_feature("expenses"):
        files.append((f"Expense_register_{period}.xlsx", expense_register_workbook(tenant, start, end).getvalue()))
    files.append((f"Profit_and_loss_{period}.xlsx",
                  pnl_workbook(statement(tenant, start, end, None, "previous"), "All locations").getvalue()))
    files.append((f"Tally_{period}.zip", tally_zip(tenant, start, end).getvalue()))
    return files


def safe_name(tenant) -> str:
    return "".join(c if c.isalnum() or c in " -_" else "_" for c in tenant.legal_name).strip() or f"tenant-{tenant.id}"


def monthly_pack_zip(tenants, period: str) -> BytesIO:
    """Several clients' packs in one ZIP, a folder per client."""
    buf = BytesIO()
    with ZipFile(buf, "w", ZIP_DEFLATED) as zf:
        for tenant in tenants:
            for name, data in pack_files(tenant, period):
                zf.writestr(f"{safe_name(tenant)}/{name}", data)
    buf.seek(0)
    return buf
