"""Workbook plumbing for bulk import/export - openpyxl only, no JS.

A template, an export and an upload all share one column layout per
ImportSpec, so "export, fix in Excel, import back" round-trips: import
updates the records it recognizes and creates the rest.
"""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.bulk.specs import ImportContext, ImportSpec, RowResult

MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_ROWS = 5000
# How far down the template pre-formats code columns as text.
_TEXT_FORMATTED_ROWS = 1000

_HEADER_FILL = PatternFill("solid", fgColor="1F7A4D")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


class WorkbookError(ValueError):
    pass


def _write_header(ws, spec: ImportSpec) -> None:
    for i, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header + (" *" if col.required else ""))
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = max(14, len(col.header) + 4)
        if col.as_text:
            # Pre-format the whole column as text so a typed 0401 keeps
            # its leading zero and a long barcode isn't turned into 8E+12.
            for row in range(2, _TEXT_FORMATTED_ROWS + 2):
                ws.cell(row=row, column=i).number_format = "@"
    ws.freeze_panes = "A2"


def _write_help_sheet(wb: Workbook, spec: ImportSpec) -> None:
    ws = wb.create_sheet("How to fill")
    ws.append(["Column", "Required", "What to enter", "Example"])
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for col in spec.columns:
        ws.append([col.header, "Yes" if col.required else "", col.help, col.example])
    ws.append([])
    ws.append(["Fill one row per record on the first sheet. Don't rename or reorder the headings."])
    ws.append(["Uploading a row that matches an existing record updates it; anything else is added."])
    for letter, width in (("A", 24), ("B", 10), ("C", 70), ("D", 22)):
        ws.column_dimensions[letter].width = width


def build_template(spec: ImportSpec) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_title
    _write_header(ws, spec)
    _write_help_sheet(wb, spec)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def build_export(spec: ImportSpec, tenant) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = spec.sheet_title
    _write_header(ws, spec)
    for r, record in enumerate(spec.export_rows(tenant), start=2):
        for c, col in enumerate(spec.columns, start=1):
            value = record.get(col.key)
            if isinstance(value, Decimal):
                value = float(value) if not col.as_text else str(value)
            ws.cell(row=r, column=c, value=value)
    _write_help_sheet(wb, spec)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _normalize(header) -> str:
    return " ".join(str(header or "").replace("*", "").split()).strip().lower()


def read_rows(spec: ImportSpec, data: bytes) -> list[tuple[int, dict]]:
    """[(sheet row number, {column key: raw cell value})] for every
    non-blank row of the first sheet."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise WorkbookError("That file is too big - 2 MB at most. Split it into smaller files.")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise WorkbookError("That isn't an Excel (.xlsx) file this app can read. Download the template and use that.")
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise WorkbookError("The sheet is empty.")

    by_header = {}
    for col in spec.columns:
        by_header[_normalize(col.header)] = col.key
        by_header[_normalize(col.key)] = col.key
    positions = {}
    for idx, title in enumerate(header):
        key = by_header.get(_normalize(title))
        if key and key not in positions:
            positions[key] = idx
    missing = [col.header for col in spec.columns if col.required and col.key not in positions]
    if missing:
        raise WorkbookError(
            "These columns are missing from the first sheet: " + ", ".join(missing) + ". Download the template and use that."
        )

    out = []
    for offset, values in enumerate(rows, start=2):
        if values is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue
        raw = {key: (values[idx] if idx < len(values) else None) for key, idx in positions.items()}
        out.append((offset, raw))
        if len(out) > MAX_ROWS:
            raise WorkbookError(f"More than {MAX_ROWS} rows - split the file into smaller ones.")
    wb.close()
    return out


def preview(spec: ImportSpec, tenant, data: bytes) -> tuple[list[RowResult], ImportContext]:
    ctx = spec.make_context(tenant)
    results = []
    for row_number, raw in read_rows(spec, data):
        row = RowResult(row_number=row_number, values={})
        spec.parse_row(raw, ctx, row)
        results.append(row)
    return results, ctx
