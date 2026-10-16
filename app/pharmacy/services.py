"""Pharmacy rules that aren't the counter's: substitutes and the register."""

import re
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.inventory import ledger
from app.models.inventory import ProductBatch, StockMovement
from app.models.pos_bill import POSBill, POSBillLine, POSBillStatus
from app.models.product import Product


def normalise(composition: str | None) -> str:
    """"Paracetamol 500 mg" and "paracetamol  500MG" are the same salt."""
    text = re.sub(r"\s+", "", (composition or "").lower())
    return text.replace("+", ",")


def substitutes(tenant, product: Product) -> list[dict]:
    """Other active items with the same composition, cheapest first, with
    stock on hand."""
    key = normalise(product.composition)
    if not key:
        return []
    candidates = [p for p in Product.query.filter(Product.tenant_id == tenant.id, Product.is_active.is_(True),
                                                  Product.id != product.id, Product.composition.isnot(None)).all()
                  if normalise(p.composition) == key]
    stock = ledger.stock_on_hand(tenant.id, [p.id for p in candidates]) if candidates else {}
    rows = [{"id": p.id, "name": p.name, "rate": float(p.sale_price), "stock": float(stock.get(p.id, 0)),
             "composition": p.composition} for p in candidates]
    return sorted(rows, key=lambda r: (r["stock"] <= 0, r["rate"]))


def register_rows(tenant, start: date, end: date, schedules=("H1", "X")) -> list[dict]:
    """Every sale of a scheduled drug in the period, with the prescription
    and the batch it came out of - the H1 / X register (or H with it)."""
    lo = datetime.combine(start, time.min).astimezone(timezone.utc)
    hi = datetime.combine(end + timedelta(days=1), time.min).astimezone(timezone.utc)
    q = (POSBillLine.query.join(POSBill, POSBill.id == POSBillLine.bill_id).join(Product, Product.id == POSBillLine.product_id)
         .filter(POSBill.tenant_id == tenant.id, POSBill.status.in_([POSBillStatus.COMPLETED, POSBillStatus.REFUNDED]),
                 POSBill.completed_at >= lo, POSBill.completed_at < hi, Product.drug_schedule.in_(schedules))
         .order_by(POSBill.completed_at, POSBillLine.sort_order))
    rows = []
    for line in q.all():
        bill = line.bill
        batches = (ProductBatch.query.join(StockMovement, StockMovement.batch_id == ProductBatch.id)
                   .filter(StockMovement.tenant_id == tenant.id, StockMovement.source_type == "pos_bill",
                           StockMovement.source_id == bill.id, StockMovement.source_line_id == line.id).all())
        rx = bill.prescription or {}
        rows.append({"date": bill.completed_at, "bill": bill.bill_number, "drug": line.description,
                     "schedule": line.product.drug_schedule, "qty": line.qty, "unit": line.unit,
                     "batch": ", ".join(sorted({b.batch_no for b in batches})) or "-",
                     "doctor": rx.get("doctor", ""), "doctor_reg_no": rx.get("doctor_reg_no", ""),
                     "patient": rx.get("patient", ""), "patient_address": rx.get("patient_address", "")})
    return rows


def register_workbook(rows, start: date, end: date) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append([f"Scheduled drug register {start.isoformat()} to {end.isoformat()}"])
    ws["A1"].font = Font(bold=True)
    head = ["Date", "Bill", "Drug", "Schedule", "Qty", "Unit", "Batch", "Prescriber", "Reg. no.", "Patient", "Address"]
    ws.append(head)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r["date"].astimezone().strftime("%d-%m-%Y %H:%M"), r["bill"], r["drug"], r["schedule"], float(r["qty"]),
                   r["unit"], r["batch"], r["doctor"], r["doctor_reg_no"], r["patient"], r["patient_address"]])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
