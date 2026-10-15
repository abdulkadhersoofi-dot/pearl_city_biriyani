"""The margin register (Rule 32(5)): every second-hand unit sold under the
margin scheme - which purchase it came from and at what price, what it
sold for, the margin and the GST on it. Excel and PDF."""

from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO

from flask import current_app, render_template
from openpyxl import Workbook
from openpyxl.styles import Font
from weasyprint import HTML

from app.models.margin import MarginAllocation
from app.utils.gst import round_money

ZERO = Decimal("0")


@dataclass
class RegisterRow:
    item: str
    serial: str | None
    purchase_date: date
    purchase_ref: str
    supplier: str
    sale_date: date
    sale_ref: str
    qty: Decimal
    purchase_value: Decimal
    sale_value: Decimal
    raw_margin: Decimal
    margin: Decimal
    taxable_value: Decimal
    cgst: Decimal
    sgst: Decimal
    igst: Decimal

    @property
    def tax(self):
        return self.cgst + self.sgst + self.igst


def register_data(tenant, start: date, end: date) -> dict:
    """Units sold in the period (by sale date), net of any that came back
    on a sales return; voided / refunded sales are left out."""
    allocations = (
        MarginAllocation.query.filter(
            MarginAllocation.tenant_id == tenant.id, MarginAllocation.released_at.is_(None),
            MarginAllocation.sold_on >= start, MarginAllocation.sold_on <= end,
        ).order_by(MarginAllocation.sold_on, MarginAllocation.id).all()
    )
    rows = []
    for a in allocations:
        qty = Decimal(a.qty) - Decimal(a.returned_qty)
        if qty <= 0:
            continue
        share = qty / Decimal(a.qty)
        scale = (lambda v: round_money(Decimal(v) * share)) if share != 1 else (lambda v: Decimal(v))
        bill = a.purchase_line.bill
        purchase_value, sale_value = scale(a.purchase_value), scale(a.sale_value)
        rows.append(RegisterRow(
            item=a.product.name, serial=a.serial.serial_no if a.serial else None,
            purchase_date=bill.bill_date, purchase_ref=f"{bill.number} ({bill.supplier_invoice_no})",
            supplier=(bill.supplier_snapshot or {}).get("name", ""), sale_date=a.sold_on,
            sale_ref=a.document_number or "", qty=qty, purchase_value=purchase_value, sale_value=sale_value,
            raw_margin=sale_value - purchase_value, margin=scale(a.margin), taxable_value=scale(a.taxable_value),
            cgst=scale(a.cgst_amount), sgst=scale(a.sgst_amount), igst=scale(a.igst_amount),
        ))
    totals = {key: sum((getattr(r, key) for r in rows), ZERO) for key in
              ("qty", "purchase_value", "sale_value", "margin", "taxable_value", "cgst", "sgst", "igst")}
    totals["tax"] = totals["cgst"] + totals["sgst"] + totals["igst"]
    totals["losses"] = sum((-r.raw_margin for r in rows if r.raw_margin < 0), ZERO)
    return {"tenant": tenant, "start": start, "end": end, "rows": rows, "totals": totals}


HEADERS = ["Item", "Serial / IMEI", "Purchase date", "Purchase bill", "Supplier", "Sale date", "Sale invoice", "Qty",
           "Purchase price", "Sale price", "Margin", "Taxable margin", "CGST", "SGST", "IGST", "GST"]


def build_register_workbook(data: dict) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Margin register"
    ws.append([f"Margin register (Rule 32(5)) - {data['tenant'].display_name}"])
    ws.append([f"{data['start'].isoformat()} to {data['end'].isoformat()}"])
    ws.append([])
    ws.append(HEADERS)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    f = float
    for r in data["rows"]:
        ws.append([r.item, r.serial or "", r.purchase_date.isoformat(), r.purchase_ref, r.supplier, r.sale_date.isoformat(),
                   r.sale_ref, f(r.qty), f(r.purchase_value), f(r.sale_value), f(r.margin), f(r.taxable_value),
                   f(r.cgst), f(r.sgst), f(r.igst), f(r.tax)])
    t = data["totals"]
    ws.append(["Total", "", "", "", "", "", "", f(t["qty"]), f(t["purchase_value"]), f(t["sale_value"]), f(t["margin"]),
               f(t["taxable_value"]), f(t["cgst"]), f(t["sgst"]), f(t["igst"]), f(t["tax"])])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.append([])
    ws.append(["A sale below its purchase price has a margin of 0 - the loss is not set off against other sales."])
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column[3:] if c.value is not None), default=8)
        ws.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 40)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def render_register_pdf(data: dict) -> bytes:
    html = render_template("reports/margin_register_pdf.html", data=data, firm_name=data["tenant"].display_name)
    return HTML(string=html, base_url=current_app.root_path).write_pdf()
