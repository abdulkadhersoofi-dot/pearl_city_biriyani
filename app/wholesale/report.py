"""Sales and collections by salesman - invoices and counter bills they made
(excluding GST, and with it), receipts they collected."""

from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.accounts import CustomerReceipt, VoucherStatus
from app.models.business import Salesman
from app.models.invoice import Invoice, InvoiceStatus
from app.models.pos_bill import POSBill, POSBillStatus

ZERO = Decimal("0")


def salesman_report(tenant, start: date, end: date) -> list[dict]:
    lo = datetime.combine(start, time.min).astimezone(timezone.utc)
    hi = datetime.combine(end + timedelta(days=1), time.min).astimezone(timezone.utc)
    rows = {}

    def row(sid):
        if sid not in rows:
            rows[sid] = {"salesman": None, "invoices": 0, "bills": 0, "taxable": ZERO, "total": ZERO, "collected": ZERO,
                         "receipts": 0}
        return rows[sid]

    for inv in Invoice.query.filter(Invoice.tenant_id == tenant.id, Invoice.status == InvoiceStatus.ISSUED,
                                    Invoice.invoice_date >= start, Invoice.invoice_date <= end,
                                    Invoice.salesman_id.isnot(None)):
        r = row(inv.salesman_id)
        r["invoices"] += 1
        r["taxable"] += inv.total_taxable_value
        r["total"] += inv.grand_total
    for bill in POSBill.query.filter(POSBill.tenant_id == tenant.id, POSBill.status == POSBillStatus.COMPLETED,
                                     POSBill.completed_at >= lo, POSBill.completed_at < hi,
                                     POSBill.salesman_id.isnot(None)):
        r = row(bill.salesman_id)
        r["bills"] += 1
        r["taxable"] += bill.total_taxable_value
        r["total"] += bill.grand_total
    for rec in CustomerReceipt.query.filter(CustomerReceipt.tenant_id == tenant.id,
                                            CustomerReceipt.status == VoucherStatus.POSTED,
                                            CustomerReceipt.voucher_date >= start, CustomerReceipt.voucher_date <= end,
                                            CustomerReceipt.salesman_id.isnot(None)):
        r = row(rec.salesman_id)
        r["receipts"] += 1
        r["collected"] += rec.amount
    names = {s.id: s for s in Salesman.query.filter(Salesman.tenant_id == tenant.id, Salesman.id.in_(list(rows) or [0]))}
    out = []
    for sid, r in rows.items():
        r["salesman"] = names.get(sid)
        out.append(r)
    return sorted(out, key=lambda r: -r["total"])


def report_workbook(rows, start: date, end: date) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Salesmen"
    ws.append([f"Sales and collections by salesman, {start.isoformat()} to {end.isoformat()}"])
    ws["A1"].font = Font(bold=True)
    ws.append(["Salesman", "Invoices", "Counter bills", "Sales excl. GST", "Sales incl. GST", "Receipts", "Collected"])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r["salesman"].name if r["salesman"] else "-", r["invoices"], r["bills"], float(r["taxable"]),
                   float(r["total"]), r["receipts"], float(r["collected"])])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
