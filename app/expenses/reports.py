"""Expense report: posted expenses over a period, by category, branch and
month, with Excel and PDF downloads. Voided expenses and drafts are left
out - a draft hasn't happened yet, and a voided one never did."""

from collections import OrderedDict
from decimal import Decimal
from io import BytesIO

from flask import current_app, render_template
from openpyxl import Workbook
from openpyxl.styles import Font
from weasyprint import HTML

from app.models.expense import PAYMENT_MODE_LABELS, Expense, ExpenseStatus

ZERO = Decimal("0")
HEAD_OFFICE = "Head office"


def _bucket():
    return {"count": 0, "amount": ZERO, "taxable": ZERO, "tax": ZERO, "itc": ZERO}


def _add(bucket, e: Expense):
    bucket["count"] += 1
    bucket["amount"] += e.amount
    bucket["taxable"] += e.taxable_value
    bucket["tax"] += e.total_tax
    bucket["itc"] += e.itc_total


def expense_report_data(tenant, start, end, branch_ids=None, category_ids=None) -> dict:
    """branch_ids None means everywhere (head office included); a branch
    selection means just those branches."""
    query = Expense.query.filter(
        Expense.tenant_id == tenant.id, Expense.status == ExpenseStatus.POSTED,
        Expense.expense_date >= start, Expense.expense_date <= end,
    )
    if branch_ids:
        query = query.filter(Expense.branch_user_id.in_(branch_ids))
    if category_ids:
        query = query.filter(Expense.category_id.in_(category_ids))
    rows = query.order_by(Expense.expense_date, Expense.id).all()

    by_category: dict[str, dict] = {}
    by_branch: dict[str, dict] = {}
    by_month: "OrderedDict[str, dict]" = OrderedDict()
    totals = _bucket()
    for e in rows:
        _add(totals, e)
        _add(by_category.setdefault(e.category.name, _bucket()), e)
        _add(by_branch.setdefault(e.branch_user.name if e.branch_user else HEAD_OFFICE, _bucket()), e)
        _add(by_month.setdefault(e.expense_date.strftime("%Y-%m"), _bucket()), e)
    return {
        "tenant": tenant,
        "start": start,
        "end": end,
        "rows": rows,
        "totals": totals,
        "by_category": sorted(by_category.items(), key=lambda kv: kv[1]["amount"], reverse=True),
        "by_branch": sorted(by_branch.items(), key=lambda kv: kv[1]["amount"], reverse=True),
        "by_month": list(by_month.items()),
        "mode_labels": PAYMENT_MODE_LABELS,
    }


def _money(v) -> float:
    return float(v or 0)


def build_expense_workbook(data: dict) -> BytesIO:
    wb = Workbook()
    bold = Font(bold=True)
    ws = wb.active
    ws.title = "Expenses"
    ws.append([f"Expenses - {data['tenant'].display_name}"])
    ws.append([f"{data['start'].isoformat()} to {data['end'].isoformat()}"])
    ws.append([])
    headers = ["Date", "Voucher", "Category", "Payee", "Branch", "Paid by", "Reference", "Supplier GSTIN",
               "Taxable", "GST", "ITC", "Amount"]
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.font = bold
    for e in data["rows"]:
        ws.append([
            e.expense_date.isoformat(), e.number, e.category.name, e.payee,
            e.branch_user.name if e.branch_user else HEAD_OFFICE, PAYMENT_MODE_LABELS[e.payment_mode],
            e.reference or e.bill_number or "", e.supplier_gstin or "",
            _money(e.taxable_value), _money(e.total_tax), _money(e.itc_total), _money(e.amount),
        ])
    t = data["totals"]
    ws.append(["", "", "", "", "", "", "", "Total", _money(t["taxable"]), _money(t["tax"]), _money(t["itc"]),
               _money(t["amount"])])
    for cell in ws[ws.max_row]:
        cell.font = bold

    for title, key in (("By category", "by_category"), ("By branch", "by_branch"), ("By month", "by_month")):
        sheet = wb.create_sheet(title)
        sheet.append([title.replace("By ", "").title(), "Expenses", "Amount", "GST", "ITC"])
        for cell in sheet[1]:
            cell.font = bold
        for label, b in data[key]:
            sheet.append([label, b["count"], _money(b["amount"]), _money(b["tax"]), _money(b["itc"])])
        sheet.append(["Total", t["count"], _money(t["amount"]), _money(t["tax"]), _money(t["itc"])])
        for cell in sheet[sheet.max_row]:
            cell.font = bold
    for sheet in wb.worksheets:
        for column in sheet.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=8)
            sheet.column_dimensions[column[0].column_letter].width = min(max(width + 2, 10), 45)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def render_expense_pdf(data: dict) -> bytes:
    html = render_template("expenses/report_pdf.html", data=data, firm_name=data["tenant"].display_name)
    return HTML(string=html, base_url=current_app.root_path).write_pdf()
