"""Excel import/export of items and customers (app.bulk)."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook, load_workbook

from app.bulk.specs import CUSTOMER_COLUMNS, PRODUCT_COLUMNS
from app.models.audit_log import AuditLog
from app.models.category import Category
from app.models.customer import Customer
from app.models.product import Product
from app.models.tenant import RegistrationType, Tenant


def _login(client, email, password):
    return client.post("/auth/login", data={"email": email, "password": password}, follow_redirects=True)


def _xlsx(columns, rows) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.append([c.header for c in columns])
    for row in rows:
        ws.append([row.get(c.key) for c in columns])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(client, kind, buf, filename="upload.xlsx"):
    return client.post(f"/bulk/{kind}/upload", data={"file": (buf, filename)}, content_type="multipart/form-data")


def test_template_has_the_columns_and_a_help_sheet(client, db, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    resp = client.get("/bulk/products/template.xlsx")
    assert resp.status_code == 200
    wb = load_workbook(BytesIO(resp.data))
    headers = [c.value for c in wb.worksheets[0][1]]
    assert headers[0] == "Item name *"
    assert "HSN/SAC code *" in headers
    assert wb.sheetnames == ["Items", "How to fill"]


def test_import_previews_then_creates_and_updates(client, db, tenant, client_admin):
    db.session.add(Product(tenant_id=tenant.id, name="Tea", hsn_or_sac_code="0902", gst_rate=Decimal("5"), unit="NOS", sale_price=Decimal("10")))
    db.session.commit()
    _login(client, client_admin.email, "ClientSecret123")

    buf = _xlsx(PRODUCT_COLUMNS, [
        {"name": "Tea", "hsn_or_sac_code": "0902", "gst_rate": 5, "sale_price": 12},  # update
        {"name": "Coffee", "hsn_or_sac_code": 901, "gst_rate": 0.05, "sale_price": "Rs. 20", "unit": "nos", "price_includes_tax": "yes", "sku": "COF"},
        {"name": "Bad rate", "hsn_or_sac_code": "2106", "gst_rate": 7, "sale_price": 5},
        {"name": "Bad unit", "hsn_or_sac_code": "2106", "gst_rate": 5, "sale_price": 5, "unit": "plate"},
        {"name": "Dup SKU", "hsn_or_sac_code": "2106", "gst_rate": 5, "sale_price": 5, "sku": "COF"},
        {"name": "No price", "hsn_or_sac_code": "2106", "gst_rate": 5},
    ])
    resp = _upload(client, "products", buf)
    assert resp.status_code == 200
    body = resp.data.decode()
    assert "1 to add" in body
    assert "1 to update" in body
    assert "4 with problems" in body
    assert "a GST rate - use one of" in body
    assert "a GST unit code (e.g. PCS" in body
    assert "is also on row 3" in body
    assert "Sale price is required" in body
    assert Product.query.count() == 1  # nothing saved by the preview

    resp = client.post("/bulk/products/confirm", follow_redirects=True)
    assert b"Items: 1 added, 1 updated." in resp.data
    assert Product.query.filter_by(name="Tea").one().sale_price == Decimal("12.00")
    coffee = Product.query.filter_by(name="Coffee").one()
    # numeric 901 lost its leading zero in Excel; 0.05 was a %-formatted 5%
    assert coffee.hsn_or_sac_code == "0901"
    assert coffee.gst_rate == Decimal("5")
    assert coffee.unit == "NOS" and coffee.price_includes_tax is True and coffee.sale_price == Decimal("20.00")
    assert AuditLog.query.filter_by(action="bulk_import_products", tenant_id=tenant.id).count() == 1


def test_import_creates_categories_only_when_the_module_is_on(client, db, tenant, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    _upload(client, "products", _xlsx(PRODUCT_COLUMNS, [{"name": "Idli", "hsn_or_sac_code": "2106", "gst_rate": 5, "sale_price": 30, "category": "Breakfast"}]))
    client.post("/bulk/products/confirm")
    assert Category.query.count() == 0
    assert Product.query.filter_by(name="Idli").one().category_id is None

    tenant.set_feature("categories", True)
    db.session.commit()
    _upload(client, "products", _xlsx(PRODUCT_COLUMNS, [{"name": "Idli", "hsn_or_sac_code": "2106", "gst_rate": 5, "sale_price": 30, "category": "Breakfast", "sub_category": "South Indian"}]))
    client.post("/bulk/products/confirm")
    idli = Product.query.filter_by(name="Idli").one()
    assert idli.category.name == "Breakfast" and idli.sub_category.name == "South Indian"


def test_hsn_rule_uses_the_clients_turnover(client, db, tenant, client_admin):
    tenant.aato = Decimal("60000000")
    db.session.commit()
    _login(client, client_admin.email, "ClientSecret123")
    resp = _upload(client, "products", _xlsx(PRODUCT_COLUMNS, [{"name": "X", "hsn_or_sac_code": "2106", "gst_rate": 5, "sale_price": 1}]))
    assert b"at least 6 digits" in resp.data


def test_export_round_trips_through_import(client, db, tenant, client_admin):
    db.session.add(Product(tenant_id=tenant.id, name="Soap", hsn_or_sac_code="34011190", gst_rate=Decimal("18"), unit="PCS", sale_price=Decimal("40"), mrp=Decimal("45"), sku="SOAP-1", price_includes_tax=True))
    db.session.commit()
    _login(client, client_admin.email, "ClientSecret123")
    exported = client.get("/bulk/products/export.xlsx")
    assert exported.status_code == 200
    wb = load_workbook(BytesIO(exported.data))
    assert wb.worksheets[0]["A2"].value == "Soap"
    resp = _upload(client, "products", BytesIO(exported.data))
    assert b"1 to update" in resp.data and b"with problems" not in resp.data


def test_import_never_touches_another_clients_items(client, db, tenant, client_admin):
    other = Tenant(legal_name="Other", registration_type=RegistrationType.REGULAR)
    db.session.add(other)
    db.session.flush()
    theirs = Product(tenant_id=other.id, name="Tea", hsn_or_sac_code="0902", gst_rate=Decimal("5"), sale_price=Decimal("10"), sku="T1")
    db.session.add(theirs)
    db.session.commit()
    _login(client, client_admin.email, "ClientSecret123")
    _upload(client, "products", _xlsx(PRODUCT_COLUMNS, [{"name": "Tea", "hsn_or_sac_code": "0902", "gst_rate": 5, "sale_price": 99, "sku": "T1"}]))
    client.post("/bulk/products/confirm")
    db.session.refresh(theirs)
    assert theirs.sale_price == Decimal("10.00")
    assert Product.query.filter_by(tenant_id=tenant.id, name="Tea").one().sale_price == Decimal("99.00")


def test_confirm_without_an_upload_does_nothing(client, db, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    resp = client.post("/bulk/products/confirm", follow_redirects=True)
    assert b"expired" in resp.data


def test_non_excel_upload_is_rejected_cleanly(client, db, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    resp = _upload(client, "products", BytesIO(b"not,a,spreadsheet"), "items.csv")
    assert resp.status_code == 302


def test_missing_required_column_is_reported(client, db, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    wb = Workbook()
    wb.active.append(["Item name", "Sale price"])
    wb.active.append(["X", 10])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = client.post("/bulk/products/upload", data={"file": (buf, "x.xlsx")}, content_type="multipart/form-data", follow_redirects=True)
    assert b"columns are missing" in resp.data


def test_customers_import_takes_state_from_the_gstin(client, db, tenant, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    _upload(client, "customers", _xlsx(CUSTOMER_COLUMNS, [
        {"name": "Ravi Traders", "gstin": "33aaaaa0000a1z5"},
        {"name": "Walk-in Mani", "state": "Kerala", "phone": 9876543210},
        {"name": "Bad", "gstin": "NIL"},
        {"name": "No state"},
    ]))
    client.post("/bulk/customers/confirm")
    ravi = Customer.query.filter_by(name="Ravi Traders").one()
    assert (ravi.gstin, ravi.state_code, ravi.state_name) == ("33AAAAA0000A1Z5", "33", "Tamil Nadu")
    mani = Customer.query.filter_by(name="Walk-in Mani").one()
    assert mani.state_code == "32" and mani.phone == "9876543210"
    assert Customer.query.filter_by(name="Bad").first() is None
    assert Customer.query.filter_by(name="No state").first() is None


def test_staff_cannot_use_bulk_import(client, db, staff):
    _login(client, staff.email, "CashierSecret123")
    assert client.get("/bulk/products/").status_code == 403
    assert client.get("/bulk/unknown/").status_code == 403


def test_unknown_kind_is_404_for_client_admin(client, db, client_admin):
    _login(client, client_admin.email, "ClientSecret123")
    assert client.get("/bulk/suppliers/").status_code == 404
