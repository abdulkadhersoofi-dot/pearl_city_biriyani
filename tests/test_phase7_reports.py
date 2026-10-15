"""Phase 7: profit & loss (COGS from the stock ledger, cash-basis fallback,
comparison, per location), item / category gross profit, daily summary,
dashboard, GSTR-1 B2CL + HSN B2B/B2C, GSTR-3B net payable, purchase and
expense registers, Tally export, filing tracker and monthly pack - with
flags and tenant isolation."""

import zipfile
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from xml.etree import ElementTree as ET

import pytest
from openpyxl import load_workbook

from app.expenses import services as es
from app.finance import filing, journal
from app.finance.dashboard import dashboard
from app.finance.pnl import profit_and_loss, statement
from app.finance.sales_analysis import daily_summary, item_gross_profit
from app.finance.tally import tally_zip
from app.inventory import ledger
from app.invoicing.services import InvoiceInput, LineInput, create_invoice
from app.models.customer import Customer
from app.models.expense import ExpenseCategory
from app.models.filing import FilingStatus, ReturnType
from app.models.invoice_series import DocumentType
from app.models.pos_bill import PaymentMode
from app.models.product import Product
from app.models.purchase import ItcEligibility, Supplier
from app.models.tenant import BillingCycle, Gstin, RegistrationType, Tenant
from app.models.user import User, UserRole
from app.notes.services import NoteInput, create_note
from app.pos.returns import create_return, returnable_lines
from app.pos.services import CartInput, CartLineInput, PaymentInput, checkout, void_bill
from app.purchases import services as ps
from app.reports.gstr import gstr1_data, gstr3b_data
from app.reports.registers import expense_register_workbook, purchase_register_workbook

D = Decimal
TODAY = date.today()
MONTH_START = TODAY.replace(day=1)
PERIOD = TODAY.strftime("%Y-%m")


def _login(client, email, password):
    return client.post("/auth/login", data={"email": email, "password": password}, follow_redirects=True)


def _on(db, tenant, *keys):
    for key in keys:
        tenant.set_feature(key, True)
    db.session.commit()


@pytest.fixture()
def fan(db, tenant):
    p = Product(tenant_id=tenant.id, name="Table fan", hsn_or_sac_code="8414", gst_rate=D("18"), unit="NOS",
                sale_price=D("200"))
    db.session.add(p)
    db.session.commit()
    return p


@pytest.fixture()
def buyer(db, tenant):
    c = Customer(tenant_id=tenant.id, name="Ravi Stores", gstin="27AAAPR1234A1Z5", state_code="27", state_name="Maharashtra")
    db.session.add(c)
    db.session.commit()
    return c


@pytest.fixture()
def vendor(db, tenant):
    s = Supplier(tenant_id=tenant.id, name="Metro Wholesale", gstin="27BBBBB1111B1Z5", state_code="27",
                 state_name="Maharashtra")
    db.session.add(s)
    db.session.commit()
    return s


def _buy(db, tenant, user, vendor, product, qty="10", rate="100", eligibility=ItcEligibility.ELIGIBLE, bill_date=None):
    bill = ps.create_purchase_bill(tenant, user, ps.PurchaseInput(
        supplier_id=vendor.id, gstin_id=tenant.gstins.first().id, supplier_invoice_no=f"S-{qty}-{rate}-{eligibility.value}",
        supplier_invoice_date=bill_date or TODAY, location_id=ledger.default_location(tenant.id).id, bill_date=bill_date,
        lines=[ps.PurchaseLineInput(product_id=product.id, description=product.name, hsn_or_sac_code=product.hsn_or_sac_code,
                                    qty=D(qty), unit="NOS", rate=D(rate), gst_rate=D("18"), itc_eligibility=eligibility)],
    ))
    db.session.commit()
    return bill


def _sell(db, tenant, user, customer, product=None, qty="4", rate="200", pos="27", on=None, gst="18"):
    inv = create_invoice(tenant, user, InvoiceInput(
        gstin_id=tenant.gstins.first().id, document_type=DocumentType.TAX_INVOICE, customer_id=customer.id,
        place_of_supply_state_code=pos, invoice_date=on or TODAY, notes=None,
        lines=[LineInput(description=product.name if product else "Service", hsn_or_sac_code=product.hsn_or_sac_code if product else "998719",
                         qty=D(qty), rate=D(rate), discount_percent=D("0"), gst_rate=D(gst),
                         product_id=product.id if product else None, unit="NOS")],
    ))
    db.session.commit()
    return inv


def _expense(db, tenant, user, amount="500", category="Rent"):
    es.ensure_categories(tenant.id)
    cat = ExpenseCategory.query.filter_by(tenant_id=tenant.id, name=category).one()
    e = es.create_expense(tenant, user, es.ExpenseInput(expense_date=TODAY, category_id=cat.id, amount=D(amount),
                                                        payee_name="Landlord"))
    db.session.commit()
    return e


def _line(pl, code):
    return next((l.amount for l in pl.lines if l.code == code), D("0"))


# --- profit & loss ---------------------------------------------------------------------------

def test_pnl_cogs_from_stock_ledger(db, tenant, client_admin, fan, buyer, vendor):
    _on(db, tenant, "inventory")
    _buy(db, tenant, client_admin, vendor, fan)               # 10 x 100 = 1000 in stock
    _sell(db, tenant, client_admin, buyer, fan)               # 4 x 200 = 800 revenue, cost 400
    _expense(db, tenant, client_admin, "500")
    pl = profit_and_loss(tenant, MONTH_START, TODAY)
    assert pl["basis"] == "accrual"
    assert pl["net_revenue"] == D("800.00")
    assert _line(pl["cost"], "purchases") == D("1000.00") and _line(pl["cost"], "closing_stock") == D("-600.00")
    assert pl["cost_total"] == D("400.00") and pl["gross_profit"] == D("400.00") and pl["gp_pct"] == D("50.0")
    assert pl["expense_total"] == D("500.00") and pl["net_profit"] == D("-100.00")


def test_pnl_credit_note_and_blocked_gst_is_cost(db, tenant, client_admin, fan, buyer, vendor):
    _on(db, tenant, "inventory")
    _buy(db, tenant, client_admin, vendor, fan, qty="10", rate="100", eligibility=ItcEligibility.BLOCKED)
    inv = _sell(db, tenant, client_admin, buyer, fan, qty="2")
    create_note(tenant, client_admin, NoteInput(inv.id, DocumentType.CREDIT_NOTE, TODAY, "Rate difference", [
        LineInput(description="Table fan", hsn_or_sac_code="8414", qty=D("1"), rate=D("50"), discount_percent=D("0"),
                  gst_rate=D("18"), unit="NOS")]))
    db.session.commit()
    pl = profit_and_loss(tenant, MONTH_START, TODAY)
    assert pl["net_revenue"] == D("350.00")  # 400 less 50
    # GST that can't be claimed is part of cost: 2 units at 118.
    assert pl["cost_total"] == D("236.00")


def test_pnl_cash_basis_without_inventory_and_comparison(db, tenant, client_admin, fan, buyer, vendor):
    _buy(db, tenant, client_admin, vendor, fan)
    _sell(db, tenant, client_admin, buyer, fan)
    last_month = MONTH_START - timedelta(days=1)
    _sell(db, tenant, client_admin, buyer, fan, qty="1", on=last_month)
    pl = profit_and_loss(tenant, MONTH_START, TODAY)
    assert pl["basis"] == "cash" and pl["cost_total"] == D("1000.00") and pl["gross_profit"] == D("-200.00")
    st = statement(tenant, MONTH_START, TODAY, compare="previous")
    rev = next(s for s in st["sections"] if s["key"] == "revenue")
    assert rev["total"] == D("800.00") and rev["prior_total"] == D("200.00")
    assert statement(tenant, MONTH_START, TODAY, compare="none")["prior"] is None


def test_pnl_per_location_and_counter_returns(db, tenant, client_admin, staff, fan, buyer, vendor):
    _on(db, tenant, "inventory")
    _buy(db, tenant, client_admin, vendor, fan)
    _sell(db, tenant, client_admin, buyer, fan, qty="1")
    cart = CartInput(lines=[CartLineInput(description="Table fan", hsn_or_sac_code="8414", qty=D("2"), rate=D("236"),
                                          gst_rate=D("18"), product_id=fan.id, unit="NOS", price_includes_tax=True)])
    bill = checkout(tenant, staff, cart, PaymentMode.CASH)
    db.session.commit()
    create_return(bill, staff, {returnable_lines(bill)[0].line.id: D("1")}, "Faulty", PaymentMode.CASH)
    db.session.commit()
    branch_loc = ledger.location_for_user(tenant.id, staff).id
    store = ledger.default_location(tenant.id).id
    at_branch = profit_and_loss(tenant, MONTH_START, TODAY, branch_loc)
    at_store = profit_and_loss(tenant, MONTH_START, TODAY, store)
    assert _line(at_branch["revenue"], "sales_counter") == D("400.00") and _line(at_branch["revenue"], "returns") == D("-200.00")
    assert _line(at_store["revenue"], "sales_invoices") == D("200.00") and _line(at_store["revenue"], "sales_counter") == 0
    total = profit_and_loss(tenant, MONTH_START, TODAY)
    assert total["net_revenue"] == at_branch["net_revenue"] + at_store["net_revenue"]


# --- item gross profit, daily summary, dashboard -------------------------------------------

def test_item_gross_profit(db, tenant, client_admin, fan, buyer, vendor):
    _on(db, tenant, "inventory")
    _buy(db, tenant, client_admin, vendor, fan)
    _sell(db, tenant, client_admin, buyer, fan)
    _sell(db, tenant, client_admin, buyer, None, qty="1", rate="300")  # typed-in service line
    data = item_gross_profit(tenant, MONTH_START, TODAY)
    fan_row = next(r for r in data["rows"] if r.name == "Table fan")
    assert (fan_row.qty, fan_row.sales, fan_row.cost, fan_row.gross_profit) == (D("4"), D("800.00"), D("400.00"), D("400.00"))
    other = next(r for r in data["rows"] if r.key == "Items not in the catalogue")
    assert other.sales == D("300.00") and other.cost == 0 and not other.has_cost
    assert data["totals"].gross_profit == D("700.00")
    assert {c.name for c in data["categories"]} == {"Uncategorised", "-"}


def test_daily_summary_modes_hours_cashiers_voids(db, tenant, staff, fan):
    tenant.pos_payment_modes = ["cash", "upi", "card"]
    db.session.commit()
    cart = lambda: CartInput(lines=[CartLineInput(description="Fan", hsn_or_sac_code="8414", qty=D("1"), rate=D("236"),  # noqa: E731
                                                  gst_rate=D("18"), product_id=fan.id, unit="NOS", price_includes_tax=True)])
    checkout(tenant, staff, cart(), [PaymentInput(PaymentMode.CASH, D("136")), PaymentInput(PaymentMode.UPI, D("100"))])
    second = checkout(tenant, staff, cart(), PaymentMode.CASH)
    db.session.commit()
    void_bill(second, staff, "Rang up twice")
    db.session.commit()
    data = daily_summary(tenant, TODAY, TODAY)
    modes = {m["mode"]: m for m in data["modes"]}
    assert modes[PaymentMode.CASH]["sales"] == D("136") and modes[PaymentMode.UPI]["sales"] == D("100")
    assert data["bill_count"] == 1 and len(data["voids"]) == 1 and data["voids"][0].void_reason == "Rang up twice"
    assert data["cashiers"][0]["name"] == "Cashier" and data["cashiers"][0]["voids"] == 1
    assert sum(h["bills"] for h in data["hours"]) == 1


def test_dashboard_tiles_follow_flags(db, tenant, client_admin, fan, buyer):
    _sell(db, tenant, client_admin, buyer, fan, qty="1")
    d = dashboard(tenant)
    assert d["sales_today"] == D("236.00") and "receivable" not in d and "low_stock" not in d
    _on(db, tenant, "receivables", "inventory")
    d = dashboard(tenant)
    assert d["receivable"] == D("236.00") and "low_stock" in d


# --- GST returns -----------------------------------------------------------------------------

def test_gstr1_b2cl_threshold_and_hsn_split(db, tenant, client_admin, buyer):
    walk_in_ka = Customer(tenant_id=tenant.id, name="Karnataka buyer", state_code="29")
    walk_in_mh = Customer(tenant_id=tenant.id, name="Local buyer", state_code="27")
    db.session.add_all([walk_in_ka, walk_in_mh])
    db.session.commit()
    big = _sell(db, tenant, client_admin, walk_in_ka, None, qty="1", rate="90000", pos="29")    # 1,06,200 inter-state
    _sell(db, tenant, client_admin, walk_in_ka, None, qty="1", rate="80000", pos="29")          # 94,400 - B2CS
    _sell(db, tenant, client_admin, walk_in_mh, None, qty="1", rate="200000", pos="27")         # intra - B2CS
    _sell(db, tenant, client_admin, buyer, None, qty="1", rate="1000")                         # B2B
    g1 = gstr1_data(tenant, tenant.gstins.first(), PERIOD)
    assert [d["doc_number"] for d in g1["b2cl"]] == [big.invoice_number]
    assert g1["b2cl"][0]["total"] == D("106200.00") and g1["b2cl"][0]["rates"][0]["igst"] == D("16200.00")
    assert sum(b.taxable_value for b in g1["b2c"]) == D("280000.00")
    assert sum(h.taxable_value for h in g1["hsn_b2b"]) == D("1000.00")
    assert sum(h.taxable_value for h in g1["hsn_b2c"]) == D("370000.00")
    assert sum(h.taxable_value for h in g1["hsn_summary"]) == D("371000.00")
    wb = load_workbook(BytesIO(__import__("app.reports.export", fromlist=["x"]).build_gstr1_workbook(g1).getvalue()))
    assert {"B2CL", "HSN B2B", "HSN B2C"} <= set(wb.sheetnames)


def test_gstr3b_net_payable_with_itc_and_rcm(db, tenant, client_admin, fan, buyer, vendor):
    gstin = tenant.gstins.first()
    assert "payment" not in gstr3b_data(tenant, gstin, PERIOD)
    _on(db, tenant, "purchases")
    _buy(db, tenant, client_admin, vendor, fan, qty="10", rate="100")          # ITC 90 + 90
    _sell(db, tenant, client_admin, buyer, fan, qty="5")                       # output 90 + 90
    _sell(db, tenant, client_admin, buyer, None, qty="1", rate="1000")         # output 90 + 90
    pay = gstr3b_data(tenant, gstin, PERIOD)["payment"]
    assert pay["output"]["cgst"] == D("180.00") and pay["utilised"]["cgst"]["cgst"] == D("90.00")
    assert pay["cash"]["cgst"] == D("90.00") and pay["cash_total"] == D("180.00")


def test_registers(db, tenant, client_admin, fan, vendor):
    _buy(db, tenant, client_admin, vendor, fan)
    _buy(db, tenant, client_admin, vendor, fan, qty="1", eligibility=ItcEligibility.BLOCKED)
    _expense(db, tenant, client_admin, "750")
    ws = load_workbook(purchase_register_workbook(tenant, MONTH_START, TODAY)).active
    values = [[c.value for c in row] for row in ws.iter_rows(min_row=5)]
    bills = [r for r in values if r[2] == "Purchase bill"]
    assert len(bills) == 2 and sorted(r[14] for r in bills) == [0.0, 180.0]  # blocked line: no ITC
    ws = load_workbook(expense_register_workbook(tenant, MONTH_START, TODAY)).active
    assert any(row[0].value == "Total" and row[15].value == 750.0 for row in ws.iter_rows())


# --- Tally ------------------------------------------------------------------------------------

def test_tally_vouchers_balance_and_parse(db, tenant, client_admin, staff, fan, buyer, vendor):
    _buy(db, tenant, client_admin, vendor, fan, qty="3", rate="99.99")
    _sell(db, tenant, client_admin, buyer, fan, qty="3", rate="133.33")
    cart = CartInput(lines=[CartLineInput(description="Fan", hsn_or_sac_code="8414", qty=D("1"), rate=D("235.55"),
                                          gst_rate=D("18"), product_id=fan.id, unit="NOS", price_includes_tax=True)])
    checkout(tenant, staff, cart, PaymentMode.CASH)
    _expense(db, tenant, client_admin, "1180")
    db.session.commit()
    ledgers, vouchers = journal.build(tenant, MONTH_START, TODAY)
    assert vouchers and all(v.debits == v.credits for v in vouchers)
    assert {"Sales", "Purchase", "Payment"} <= {v.vch_type for v in vouchers}
    z = zipfile.ZipFile(tally_zip(tenant, MONTH_START, TODAY))
    root = ET.fromstring(z.read("2_vouchers.xml"))
    sale = next(v for v in root.iter("VOUCHER") if v.get("VCHTYPE") == "Sales" and v.findtext("PARTYLEDGERNAME") == "Ravi Stores")
    party = next(e for e in sale.iter("ALLLEDGERENTRIES.LIST") if e.findtext("LEDGERNAME") == "Ravi Stores")
    assert party.findtext("ISDEEMEDPOSITIVE") == "Yes" and party.findtext("AMOUNT").startswith("-")
    total = sum(D(e.findtext("AMOUNT")) for e in sale.iter("ALLLEDGERENTRIES.LIST"))
    assert total == 0
    masters = ET.fromstring(z.read("1_masters.xml"))
    names = {l.get("NAME"): l.findtext("PARENT") for l in masters.iter("LEDGER")}
    assert names["Ravi Stores"] == "Sundry Debtors" and names["Metro Wholesale"] == "Sundry Creditors"
    assert names["Output CGST"] == "Duties & Taxes"


# --- filing tracker, monthly pack, routes, isolation -------------------------------------

def test_due_dates():
    assert filing.due_date("2026-09", ReturnType.GSTR1) == date(2026, 10, 11)
    assert filing.due_date("2026-12", ReturnType.GSTR3B) == date(2027, 1, 20)


@pytest.fixture()
def other(db, super_admin):
    t = Tenant(legal_name="Other Co", registration_type=RegistrationType.REGULAR, onboarded_by_id=super_admin.id,
               billing_cycle=BillingCycle.MONTHLY)
    db.session.add(t)
    db.session.flush()
    db.session.add(Gstin(tenant_id=t.id, gstin="29BBBBB0000B1Z5", state_code="29", state_name="Karnataka", is_primary=True))
    admin = User(tenant_id=t.id, name="Other", email="other@x.example.com", role=UserRole.CLIENT_ADMIN, must_change_password=False)
    admin.set_password("OtherSecret123")
    db.session.add(admin)
    t.set_feature("business_reports", True)
    db.session.commit()
    return {"tenant": t, "admin": admin}


def test_flags_off_404_and_landing_unchanged(db, client, tenant, client_admin):
    _login(client, "admin@acme.example.com", "ClientSecret123")
    assert client.get("/").location.endswith("/invoices/")
    for url in ("/reports/dashboard", "/reports/profit-loss", "/reports/item-profit", "/reports/daily-summary",
                "/reports/filing", "/reports/purchase-register.xlsx", "/reports/expense-register.xlsx"):
        assert client.get(url).status_code == 404, url
    _on(db, tenant, "business_reports")
    assert client.get("/").location.endswith("/reports/dashboard")
    client.get("/auth/logout")
    landed = client.post("/auth/login", data={"email": "admin@acme.example.com", "password": "ClientSecret123"})
    assert landed.location.endswith("/reports/dashboard")


def test_report_pages_render(db, client, tenant, client_admin, staff, fan, buyer, vendor):
    _on(db, tenant, "business_reports", "inventory", "purchases", "expenses", "receivables")
    _buy(db, tenant, client_admin, vendor, fan)
    _sell(db, tenant, client_admin, buyer, fan)
    _login(client, "cashier@acme.example.com", "CashierSecret123")
    assert client.get("/reports/profit-loss").status_code == 403
    client.get("/auth/logout")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    for url in ("/reports/dashboard", "/reports/profit-loss", "/reports/profit-loss?compare=last_year&mode=range&start=2026-01-01&end=2026-12-31",
                "/reports/item-profit", "/reports/item-profit?view=categories", "/reports/daily-summary", "/reports/filing", "/reports/"):
        assert client.get(url).status_code == 200, url
    for url in ("/reports/profit-loss.xlsx", "/reports/profit-loss.pdf", "/reports/item-profit.xlsx", "/reports/daily-summary.pdf",
                "/reports/purchase-register.xlsx", "/reports/expense-register.xlsx"):
        r = client.get(url)
        assert r.status_code == 200 and len(r.data) > 1000, url
    assert b"Gross profit" in client.get("/reports/profit-loss").data


def test_filing_tracker_flow_and_isolation(db, client, tenant, client_admin, super_admin, other):
    _on(db, tenant, "business_reports")
    gstin = tenant.gstins.first()
    _login(client, "admin@acme.example.com", "ClientSecret123")
    client.post("/reports/filing", data={"gstin_id": gstin.id, "period": "2026-09", "action": "ready"})
    rows = FilingStatus.query.filter_by(tenant_id=tenant.id, period="2026-09").all()
    assert len(rows) == 2 and all(r.data_ready for r in rows)
    # Another client can't touch this GSTIN or see this status.
    client.get("/auth/logout")
    _login(client, "other@x.example.com", "OtherSecret123")
    assert client.post("/reports/filing", data={"gstin_id": gstin.id, "period": "2026-09", "action": "not_ready"}).status_code == 404
    assert gstin.gstin.encode() not in client.get("/reports/filing?period=2026-09").data
    assert all(r.data_ready for r in FilingStatus.query.filter_by(tenant_id=tenant.id).all())
    assert client.get("/admin/filing").status_code == 403
    client.get("/auth/logout")
    # The firm marks it filed.
    _login(client, super_admin.email, "SuperSecret123")
    page = client.get("/admin/filing?period=2026-09")
    assert page.status_code == 200 and b"Acme Traders" in page.data and b"11 Oct" in page.data
    client.post(f"/admin/clients/{tenant.id}/filing", data={"gstin_id": gstin.id, "period": "2026-09",
                                                             "return_type": "gstr1", "action": "filed",
                                                             "filed_on": "2026-10-10", "arn": "AA270926123456X"})
    row = FilingStatus.query.filter_by(tenant_id=tenant.id, period="2026-09", return_type=ReturnType.GSTR1).one()
    assert row.filed and row.filed_on == date(2026, 10, 10) and row.arn == "AA270926123456X"
    # A GSTIN of another client can't be posted through this client's URL.
    other_gstin = other["tenant"].gstins.first()
    assert client.post(f"/admin/clients/{tenant.id}/filing", data={"gstin_id": other_gstin.id, "period": "2026-09",
                                                                    "action": "ready"}).status_code == 404


def test_auditor_scope_on_tracker_pack_and_tally(db, client, tenant, super_admin, other):
    auditor = User(name="Auditor A", email="aud@example.com", role=UserRole.AUDITOR, must_change_password=False,
                   billing_cycle=BillingCycle.MONTHLY)
    auditor.set_password("AuditSecret123")
    db.session.add(auditor)
    db.session.commit()
    other["tenant"].auditor_id = auditor.id
    db.session.commit()
    _login(client, "aud@example.com", "AuditSecret123")
    page = client.get("/admin/filing")
    assert b"Other Co" in page.data and b"Acme Traders" not in page.data
    assert client.get(f"/admin/clients/{tenant.id}/monthly-pack.zip?period={PERIOD}").status_code == 403
    assert client.get(f"/admin/clients/{tenant.id}/tally.zip?period={PERIOD}").status_code == 403
    assert client.post(f"/admin/clients/{tenant.id}/filing", data={"gstin_id": tenant.gstins.first().id,
                                                                    "period": PERIOD, "action": "ready"}).status_code == 403
    r = client.get(f"/admin/clients/{other['tenant'].id}/monthly-pack.zip?period={PERIOD}")
    assert r.status_code == 200
    names = zipfile.ZipFile(BytesIO(r.data)).namelist()
    assert any(n.endswith(f"GSTR1_{PERIOD}.json") for n in names) and f"Profit_and_loss_{PERIOD}.xlsx" in names
    assert f"Tally_{PERIOD}.zip" in names and any("HSN_summary" in n for n in names)


def test_monthly_pack_registers_follow_flags(db, tenant, client_admin):
    names = [n for n, _ in filing.pack_files(tenant, PERIOD)]
    assert not any("register" in n for n in names)
    _on(db, tenant, "purchases", "expenses")
    names = [n for n, _ in filing.pack_files(tenant, PERIOD)]
    assert f"Purchase_register_{PERIOD}.xlsx" in names and f"Expense_register_{PERIOD}.xlsx" in names
