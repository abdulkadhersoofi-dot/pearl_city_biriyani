"""Phase 4 expenses: GST on expense bills (one tax module), voucher
numbering, drafts, voids, recurring expenses, petty cash from the POS
drawer, the expense report, expense ITC in GSTR-3B, the module flag, and
tenant isolation for every new table."""

import json
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.expenses import services as es
from app.expenses.reports import expense_report_data
from app.models.audit_log import AuditLog
from app.models.expense import (
    DEFAULT_CATEGORIES,
    Expense,
    ExpenseCategory,
    ExpensePaymentMode,
    ExpenseStatus,
    Frequency,
    RecurringExpense,
)
from app.models.media import UploadedImage
from app.models.purchase import ItcEligibility
from app.models.tenant import Gstin, RegistrationType, Tenant
from app.models.user import User, UserRole
from app.reports.gstr import gstr3b_data
from app.utils import itc
from app.utils.tenant_scope import tenant_query

D = Decimal
E, I, B = ItcEligibility.ELIGIBLE, ItcEligibility.INELIGIBLE, ItcEligibility.BLOCKED
TODAY = date.today()


def _login(client, email, password):
    return client.post("/auth/login", data={"email": email, "password": password}, follow_redirects=True)


def _on(db, tenant, *keys):
    for key in keys:
        tenant.set_feature(key, True)
    db.session.commit()


def _cat(tenant, name="Rent"):
    es.ensure_categories(tenant.id)
    return ExpenseCategory.query.filter_by(tenant_id=tenant.id, name=name).one()


def _input(tenant, amount="1000", category="Rent", **kw):
    return es.ExpenseInput(expense_date=kw.pop("expense_date", TODAY), category_id=_cat(tenant, category).id,
                           amount=D(amount), payee_name=kw.pop("payee_name", "Landlord"), **kw)


# --- the GST math (app.utils.itc) -----------------------------------------------------------

def test_inclusive_bill_total_backs_out_tax_intra_state():
    g = itc.compute_expense_gst(D("11800"), D("18"), True, True, False, E, True)
    assert (g["taxable_value"], g["cgst_amount"], g["sgst_amount"], g["igst_amount"]) == (D("10000.00"), D("900.00"), D("900.00"), D("0.00"))
    assert g["amount"] == D("11800.00") and g["itc_cgst"] == D("900.00")


def test_exclusive_amount_inter_state_is_igst():
    g = itc.compute_expense_gst(D("5000"), D("12"), False, False, False, E, True)
    assert (g["taxable_value"], g["igst_amount"], g["amount"]) == (D("5000.00"), D("600.00"), D("5600.00"))


def test_zero_rated_expense_bill():
    g = itc.compute_expense_gst(D("750"), D("0"), True, True, False, E, True)
    assert g["taxable_value"] == D("750.00") and g["cgst_amount"] + g["sgst_amount"] + g["igst_amount"] == 0
    assert g["amount"] == D("750.00")


def test_rounding_half_up_on_odd_paise():
    g = itc.compute_expense_gst(D("100.05"), D("5"), True, False, False, E, True)
    # 5% of 100.05 = 5.0025 -> 5.00, split 2.50 + 2.50
    assert g["cgst_amount"] + g["sgst_amount"] == D("5.00")
    g = itc.compute_expense_gst(D("33.33"), D("18"), True, False, False, E, True)
    assert g["cgst_amount"] + g["sgst_amount"] == D("6.00")  # 5.9994 -> 6.00


def test_rcm_pays_payee_only_taxable_and_ignores_inclusive():
    g = itc.compute_expense_gst(D("10000"), D("18"), True, True, True, E, True)
    assert g["taxable_value"] == D("10000.00") and g["amount"] == D("10000.00")
    assert g["cgst_amount"] == D("900.00") and g["itc_cgst"] == D("900.00")


def test_blocked_or_unclaimable_gives_no_credit():
    assert itc.compute_expense_gst(D("1180"), D("18"), True, True, False, B, True)["itc_cgst"] == 0
    assert itc.compute_expense_gst(D("1180"), D("18"), True, True, False, E, False)["itc_sgst"] == 0


# --- vouchers -------------------------------------------------------------------------------

def test_categories_seeded_once_and_editable(db, tenant):
    cats = es.ensure_categories(tenant.id)
    assert sorted(c.name for c in cats) == sorted(DEFAULT_CATEGORIES)
    es.ensure_categories(tenant.id)
    assert ExpenseCategory.query.filter_by(tenant_id=tenant.id).count() == len(DEFAULT_CATEGORIES)
    es.add_category(tenant.id, "Staff welfare")
    with pytest.raises(es.ExpenseError, match="already"):
        es.add_category(tenant.id, "staff WELFARE")


def test_posted_expense_is_numbered_and_drafts_never_take_numbers(db, tenant, client_admin):
    draft = es.create_expense(tenant, client_admin, _input(tenant), post=False)
    a = es.create_expense(tenant, client_admin, _input(tenant))
    b = es.create_expense(tenant, client_admin, _input(tenant, "250", "Fuel", payment_mode=ExpensePaymentMode.UPI))
    assert draft.number is None and draft.status == ExpenseStatus.DRAFT
    assert a.number.startswith("EXP/") and a.number != b.number
    es.post_expense(draft, client_admin)
    assert draft.number.endswith("0003")
    with pytest.raises(es.ExpenseError, match="already posted"):
        es.post_expense(draft, client_admin)


def test_validation(db, tenant, client_admin):
    with pytest.raises(es.ExpenseError, match="more than 0"):
        es.create_expense(tenant, client_admin, _input(tenant, "0"))
    with pytest.raises(es.ExpenseError, match="future"):
        es.create_expense(tenant, client_admin, _input(tenant, expense_date=TODAY + timedelta(days=1)))
    with pytest.raises(es.ExpenseError, match="who was paid"):
        es.create_expense(tenant, client_admin, _input(tenant, payee_name=""))
    with pytest.raises(es.ExpenseError, match="GST rate"):
        es.create_expense(tenant, client_admin, _input(tenant, is_gst_bill=True, supplier_gstin="27AAAAA1111A1Z5", gst_rate=D("7")))
    with pytest.raises(es.ExpenseError, match="GSTIN"):
        es.create_expense(tenant, client_admin, _input(tenant, is_gst_bill=True, gst_rate=D("18")))


def test_gst_bill_expense_intra_and_inter_state(db, tenant, client_admin):
    # The tenant fixture's GSTIN is Maharashtra (27).
    intra = es.create_expense(tenant, client_admin, _input(tenant, "11800", is_gst_bill=True, supplier_gstin="27aaaaa1111a1z5",
                                                           gst_rate=D("18"), amount_includes_gst=True))
    assert intra.supplier_gstin == "27AAAAA1111A1Z5"
    assert (intra.taxable_value, intra.cgst_amount, intra.itc_total, intra.amount) == (D("10000.00"), D("900.00"), D("1800.00"), D("11800.00"))
    inter = es.create_expense(tenant, client_admin, _input(tenant, "1000", "Telephone/Internet", is_gst_bill=True,
                                                           supplier_gstin="29AAAAA1111A1Z5", gst_rate=D("18"),
                                                           amount_includes_gst=False))
    assert (inter.igst_amount, inter.cgst_amount, inter.amount) == (D("180.00"), D("0.00"), D("1180.00"))


@pytest.mark.parametrize("reg", ["composition", "unregistered"])
def test_no_itc_on_expenses_for_composition_or_unregistered(db, tenant, client_admin, reg):
    tenant.registration_type = RegistrationType(reg)
    db.session.commit()
    e = es.create_expense(tenant, client_admin, _input(tenant, "1180", is_gst_bill=True, supplier_gstin="27AAAAA1111A1Z5",
                                                       gst_rate=D("18")))
    assert e.total_tax == D("180.00") and e.itc_total == 0 and e.itc_eligibility == I


def test_void_keeps_the_voucher(db, tenant, client_admin):
    e = es.create_expense(tenant, client_admin, _input(tenant))
    with pytest.raises(es.ExpenseError, match="reason"):
        es.void_expense(e, client_admin, " ")
    es.void_expense(e, client_admin, "entered twice")
    assert e.status == ExpenseStatus.VOIDED and e.number and Expense.query.count() == 1
    with pytest.raises(es.ExpenseError):
        es.void_expense(e, client_admin, "again")


# --- recurring -------------------------------------------------------------------------------

def test_monthly_due_dates_keep_the_start_day():
    d = date(2026, 1, 31)
    seq = []
    for _ in range(4):
        d = es.next_due(d, Frequency.MONTHLY, 31)
        seq.append(d)
    assert seq == [date(2026, 2, 28), date(2026, 3, 31), date(2026, 4, 30), date(2026, 5, 31)]
    assert es.next_due(date(2026, 1, 15), Frequency.QUARTERLY, 15) == date(2026, 4, 15)
    assert es.next_due(date(2026, 1, 15), Frequency.WEEKLY, 15) == date(2026, 1, 22)
    assert es.next_due(date(2024, 2, 29), Frequency.YEARLY, 29) == date(2025, 2, 28)


def test_recurring_makes_one_draft_per_period_idempotently(db, tenant, client_admin):
    rec = es.save_recurring(tenant, client_admin, es.RecurringInput(
        name="Shop rent", category_id=_cat(tenant).id, amount=D("25000"), frequency=Frequency.MONTHLY,
        start_date=date(2026, 1, 5), payee_name="Mr. Kumar"))
    made = es.generate_due_drafts(tenant, today=date(2026, 3, 10))
    assert [d.due_date for d in made] == [date(2026, 1, 5), date(2026, 2, 5), date(2026, 3, 5)]
    assert all(d.status == ExpenseStatus.DRAFT and d.number is None and d.amount == D("25000") for d in made)
    assert es.generate_due_drafts(tenant, today=date(2026, 3, 10)) == []
    assert rec.next_due_date == date(2026, 4, 5)
    # Skipping a draft doesn't bring it back.
    es.void_expense(made[0], client_admin, "paid in advance last year")
    assert es.generate_due_drafts(tenant, today=date(2026, 3, 31)) == []
    db.session.commit()


def test_recurring_stops_after_its_end_date(db, tenant, client_admin):
    rec = es.save_recurring(tenant, client_admin, es.RecurringInput(
        name="Hoarding", category_id=_cat(tenant, "Advertisement").id, amount=D("3000"), frequency=Frequency.MONTHLY,
        start_date=date(2026, 1, 1), end_date=date(2026, 2, 15), payee_name="Ads Co"))
    assert len(es.generate_due_drafts(tenant, today=date(2026, 6, 1))) == 2
    assert not rec.is_active


def test_draft_from_recurring_can_get_gst_details_then_post(db, tenant, client_admin):
    es.save_recurring(tenant, client_admin, es.RecurringInput(
        name="Rent", category_id=_cat(tenant).id, amount=D("11800"), frequency=Frequency.MONTHLY,
        start_date=TODAY - timedelta(days=1), payee_name="Landlord"))
    draft = es.generate_due_drafts(tenant)[0]
    es.update_draft(draft, tenant, _input(tenant, "11800", expense_date=draft.expense_date, is_gst_bill=True,
                                          supplier_gstin="27AAAAA1111A1Z5", gst_rate=D("18")))
    es.post_expense(draft, client_admin)
    assert draft.number and draft.itc_total == D("1800.00") and draft.recurring_id


# --- petty cash ------------------------------------------------------------------------------

def _open_day(db, tenant, user):
    from app.pos.sessions import open_session

    _on(db, tenant, "day_close")
    open_session(tenant, user, D("500"))
    db.session.commit()


def test_petty_cash_cash_out_books_an_expense_only_with_the_module(client, db, tenant, staff):
    _open_day(db, tenant, staff)
    _login(client, "cashier@acme.example.com", "CashierSecret123")
    client.post("/pos/day/cash", data={"kind": "out", "amount": "120", "reason": "Tea"})
    assert Expense.query.count() == 0  # expenses off: exactly as before

    _on(db, tenant, "expenses")
    page = client.get("/pos/day/close").data.decode()
    assert 'name="expense_category_id"' in page
    client.post("/pos/day/cash", data={"kind": "out", "amount": "80.50", "reason": "Courier"})
    e = Expense.query.one()
    assert (e.amount, e.payment_mode, e.status, e.branch_user_id, e.payee_name) == (
        D("80.50"), ExpensePaymentMode.CASH, ExpenseStatus.POSTED, staff.id, "Courier")
    assert e.category.name == "Miscellaneous" and e.cash_movement_id is not None

    fuel = _cat(tenant, "Fuel")
    db.session.commit()
    client.post("/pos/day/cash", data={"kind": "out", "amount": "300", "reason": "Petrol", "expense_category_id": str(fuel.id)})
    client.post("/pos/day/cash", data={"kind": "out", "amount": "5000", "reason": "To bank", "expense_category_id": "none"})
    client.post("/pos/day/cash", data={"kind": "in", "amount": "200", "reason": "Float"})
    rows = Expense.query.order_by(Expense.id).all()
    assert [(r.payee_name, r.category.name) for r in rows] == [("Courier", "Miscellaneous"), ("Petrol", "Fuel")]


# --- screens, flag, attachment ----------------------------------------------------------------

PAGES = ["/expenses/", "/expenses/new", "/expenses/categories", "/expenses/recurring", "/expenses/recurring/new",
         "/expenses/report", "/expenses/report.xlsx", "/expenses/report.pdf", "/expenses/?status=draft"]


def test_expenses_404_and_hidden_when_off(client, db, tenant, client_admin):
    _login(client, "admin@acme.example.com", "ClientSecret123")
    for url in PAGES:
        assert client.get(url).status_code == 404, url
    assert b'href="/expenses/"' not in client.get("/").data


def test_staff_cannot_open_expenses(client, db, tenant, staff):
    _on(db, tenant, "expenses")
    _login(client, "cashier@acme.example.com", "CashierSecret123")
    assert client.get("/expenses/").status_code == 403


def test_screens_render_and_an_expense_posts_through_the_form(client, db, tenant, client_admin):
    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    for url in PAGES:
        assert client.get(url).status_code == 200, url
    assert b'href="/expenses/"' in client.get("/expenses/").data
    rent = _cat(tenant)
    db.session.commit()
    pdf = b"%PDF-1.4 electricity bill"
    resp = client.post("/expenses/new", data={
        "expense_date": TODAY.isoformat(), "category_id": str(rent.id), "payee_name": "Landlord", "amount": "11800",
        "payment_mode": "bank", "reference": "UTR123", "is_gst_bill": "1", "supplier_gstin": "27AAAAA1111A1Z5",
        "bill_number": "R/7", "gst_rate": "18", "amount_includes_gst": "1", "itc_eligibility": "eligible",
        "action": "post", "attachment": (BytesIO(pdf), "bill.pdf"),
    }, content_type="multipart/form-data")
    assert resp.status_code == 302
    e = Expense.query.one()
    assert e.number and e.itc_total == D("1800.00") and e.reference == "UTR123"
    assert AuditLog.query.filter_by(action="expense_posted", entity_id=e.id).count() == 1
    assert client.get(f"/expenses/{e.id}").status_code == 200
    image = db.session.get(UploadedImage, e.attachment_id)
    assert image.is_private and client.get(f"/media/{image.id}").status_code == 404
    att = client.get(f"/expenses/{e.id}/attachment")
    assert att.status_code == 200 and att.data == pdf
    # A refused form keeps what was typed.
    resp = client.post("/expenses/new", data={"expense_date": TODAY.isoformat(), "category_id": str(rent.id),
                                               "payee_name": "Typo Traders", "amount": "abc"})
    assert resp.status_code == 400 and b"Typo Traders" in resp.data
    resp = client.post(f"/expenses/{e.id}/void", data={"reason": "duplicate"})
    assert resp.status_code == 302 and e.status == ExpenseStatus.VOIDED
    assert client.get(f"/expenses/{e.id}/edit").status_code == 302  # posted/voided can't be edited


def test_draft_edit_and_post_through_the_screens(client, db, tenant, client_admin):
    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    rent = _cat(tenant)
    db.session.commit()
    resp = client.post("/expenses/recurring/new", data={
        "name": "Rent", "category_id": str(rent.id), "amount": "20000", "frequency": "monthly",
        "start_date": (TODAY - timedelta(days=40)).isoformat(), "payee_name": "Landlord", "payment_mode": "bank"})
    assert resp.status_code == 302
    page = client.get("/expenses/").data
    assert b"draft expense" in page
    draft = Expense.query.filter_by(status=ExpenseStatus.DRAFT).order_by(Expense.due_date).first()
    assert client.get(f"/expenses/{draft.id}/edit").status_code == 200
    resp = client.post(f"/expenses/{draft.id}/edit", data={
        "expense_date": draft.expense_date.isoformat(), "category_id": str(rent.id), "payee_name": "Landlord",
        "amount": "20500", "payment_mode": "cheque", "action": "post"})
    assert resp.status_code == 302 and draft.number and draft.amount == D("20500.00")
    rec = RecurringExpense.query.one()
    client.post(f"/expenses/recurring/{rec.id}/toggle")
    assert not rec.is_active


# --- report + GSTR-3B ---------------------------------------------------------------------------

def test_report_groups_by_category_branch_and_month(client, db, tenant, client_admin, staff):
    es.create_expense(tenant, client_admin, _input(tenant, "1000"))
    es.create_expense(tenant, client_admin, _input(tenant, "250", "Fuel", branch_user_id=staff.id))
    es.create_expense(tenant, client_admin, _input(tenant, "999"), post=False)  # draft: left out
    voided = es.create_expense(tenant, client_admin, _input(tenant, "777"))
    es.void_expense(voided, client_admin, "x")
    db.session.commit()
    data = expense_report_data(tenant, TODAY - timedelta(days=1), TODAY)
    assert data["totals"]["amount"] == D("1250.00") and data["totals"]["count"] == 2
    assert dict(data["by_category"])["Fuel"]["amount"] == D("250.00")
    assert set(dict(data["by_branch"])) == {"Head office", "Cashier"}
    only_branch = expense_report_data(tenant, TODAY, TODAY, branch_ids=[staff.id])
    assert only_branch["totals"]["amount"] == D("250.00")

    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    wb = load_workbook(BytesIO(client.get("/expenses/report.xlsx?mode=day").data))
    assert wb.sheetnames == ["Expenses", "By category", "By branch", "By month"]
    assert client.get("/expenses/report.pdf?mode=day").data.startswith(b"%PDF")
    fuel = _cat(tenant, "Fuel")
    page = client.get(f"/expenses/report?mode=day&category_id={fuel.id}").data.decode()
    assert "Rs. 250.00" in page and "All categories" in page


def test_expense_itc_flows_into_gstr3b(db, tenant, client_admin):
    gstin = tenant.gstins.first()
    period = TODAY.strftime("%Y-%m")
    es.create_expense(tenant, client_admin, _input(tenant, "11800", is_gst_bill=True, supplier_gstin="27AAAAA1111A1Z5",
                                                   gst_rate=D("18")))                      # 900 + 900 eligible
    es.create_expense(tenant, client_admin, _input(tenant, "1180", "Office Expenses", is_gst_bill=True,
                                                   supplier_gstin="27AAAAA1111A1Z5", gst_rate=D("18"),
                                                   itc_eligibility=B))                     # 90 + 90 blocked
    es.create_expense(tenant, client_admin, _input(tenant, "5000", "Professional Fees", is_gst_bill=True, is_rcm=True,
                                                   gst_rate=D("18")))                      # RCM 450 + 450
    es.create_expense(tenant, client_admin, _input(tenant, "500", "Fuel"))                 # no GST bill
    v = es.create_expense(tenant, client_admin, _input(tenant, "11800", is_gst_bill=True,
                                                       supplier_gstin="27AAAAA1111A1Z5", gst_rate=D("18")))
    es.void_expense(v, client_admin, "x")                                                  # voided: left out
    db.session.commit()
    assert "itc" not in gstr3b_data(tenant, gstin, period)  # module off: 3B unchanged

    _on(db, tenant, "expenses")
    data = gstr3b_data(tenant, gstin, period)["itc"]
    assert data["rcm_inward"]["taxable_value"] == D("5000.00") and data["rcm_inward"]["cgst"] == D("450.00")
    assert data["a3_rcm"]["cgst"] == D("450.00")
    assert data["a5_other"]["cgst"] == D("990.00")
    assert data["b1_ineligible"]["cgst"] == D("90.00")
    assert data["net"]["cgst"] == D("1350.00")  # 450 + 990 - 90
    assert data["expense_count"] == 3


def test_gstr3b_page_with_only_expenses_on(client, db, tenant, client_admin):
    _on(db, tenant, "expenses")
    es.create_expense(tenant, client_admin, _input(tenant, "1180", is_gst_bill=True, supplier_gstin="27AAAAA1111A1Z5",
                                                   gst_rate=D("18")))
    db.session.commit()
    _login(client, "admin@acme.example.com", "ClientSecret123")
    q = f"gstin_id={tenant.gstins.first().id}&period={TODAY.strftime('%Y-%m')}"
    body = client.get(f"/reports/gstr3b?{q}").data.decode()
    assert "All other ITC" in body and "1 expense(s) with a GST bill" in body
    assert "GSTR-2B match" not in body  # purchases module is off
    assert json.loads(client.get(f"/reports/gstr3b.json?{q}").data)["itc"]["a5_other"]["cgst"] in ("90.00", 90.0, "90")


# --- tenant isolation -----------------------------------------------------------------------------

@pytest.fixture()
def other(db):
    b = Tenant(legal_name="Beta Corp", registration_type=RegistrationType.REGULAR, features={"expenses": True})
    db.session.add(b)
    db.session.flush()
    db.session.add(Gstin(tenant_id=b.id, gstin="29BBBBB0000B1Z5", state_code="29", state_name="Karnataka", is_primary=True))
    admin = User(tenant_id=b.id, name="Beta Admin", email="admin@beta.example.com", role=UserRole.CLIENT_ADMIN,
                 must_change_password=False)
    admin.set_password("BetaSecret123")
    branch = User(tenant_id=b.id, name="Beta Branch", email="branch@beta.example.com", role=UserRole.STAFF,
                  must_change_password=False)
    branch.set_password("BetaSecret123")
    db.session.add_all([admin, branch])
    db.session.commit()
    expense = es.create_expense(b, admin, _input(b, "4321", payee_name="Beta Landlord"))
    att = UploadedImage(content_type="application/pdf", data=b"%PDF secret", is_private=True)
    db.session.add(att)
    db.session.flush()
    expense.attachment_id = att.id
    rec = es.save_recurring(b, admin, es.RecurringInput(name="Beta rent", category_id=_cat(b).id, amount=D("1"),
                                                        frequency=Frequency.MONTHLY, start_date=TODAY, payee_name="x"))
    db.session.commit()
    return {"tenant": b, "admin": admin, "branch": branch, "expense": expense, "recurring": rec,
            "category": _cat(b)}


def test_tenant_a_never_reads_or_writes_tenant_b_expenses(client, db, tenant, client_admin, other):
    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    e, rec, cat = other["expense"], other["recurring"], other["category"]
    for url in (f"/expenses/{e.id}", f"/expenses/{e.id}/edit", f"/expenses/{e.id}/attachment",
                f"/expenses/recurring/{rec.id}", f"/media/{e.attachment_id}"):
        assert client.get(url).status_code == 404, url
    for url, data in ((f"/expenses/{e.id}/void", {"reason": "x"}), (f"/expenses/{e.id}/post", {}),
                      (f"/expenses/recurring/{rec.id}/toggle", {}), (f"/expenses/categories/{cat.id}", {"action": "off"})):
        assert client.post(url, data=data).status_code == 404, url
    for url in ("/expenses/", "/expenses/?status=voided", "/expenses/report", "/expenses/categories", "/expenses/recurring"):
        assert b"Beta" not in client.get(url).data, url
    assert e.status == ExpenseStatus.POSTED and rec.is_active and cat.is_active
    # B's category / branch smuggled into A's form is refused.
    with pytest.raises(es.ExpenseError, match="category"):
        es.create_expense(tenant, client_admin, es.ExpenseInput(expense_date=TODAY, category_id=cat.id, amount=D("1"),
                                                                payee_name="x"))
    with pytest.raises(es.ExpenseError, match="branch"):
        es.create_expense(tenant, client_admin, _input(tenant, branch_user_id=other["branch"].id))
    with pytest.raises(es.ExpenseError, match="registrations"):
        es.create_expense(tenant, client_admin, _input(tenant, gstin_id=other["tenant"].gstins.first().id))
    db.session.rollback()
    # Drafts are only ever made for a client's own recurring expenses.
    assert es.generate_due_drafts(tenant, today=TODAY + timedelta(days=400)) == []


def test_tenant_query_scopes_expense_tables(app, db, tenant, client_admin, other):
    from flask_login import login_user

    es.create_expense(tenant, client_admin, _input(tenant))
    es.save_recurring(tenant, client_admin, es.RecurringInput(name="Rent", category_id=_cat(tenant).id, amount=D("1"),
                                                              frequency=Frequency.MONTHLY, start_date=TODAY, payee_name="x"))
    db.session.commit()
    with app.test_request_context():
        login_user(client_admin)
        for model in (Expense, ExpenseCategory, RecurringExpense):
            rows = tenant_query(model).all()
            assert rows and all(r.tenant_id == tenant.id for r in rows), model.__name__


def test_alpine_attributes_survive_html_parsing(client, db, tenant, client_admin):
    from html.parser import HTMLParser

    class Collect(HTMLParser):
        def __init__(self):
            super().__init__()
            self.values = []

        def handle_starttag(self, tag, attrs):
            self.values += [v for k, v in attrs if (k.startswith("x-") or k.startswith(":") or k.startswith("@")) and v]

    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    for url in ("/expenses/new", "/expenses/report"):
        parser = Collect()
        parser.feed(client.get(url).data.decode())
        assert parser.values, url
        for value in parser.values:
            assert value.count("(") == value.count(")") and value.count("{") == value.count("}"), (url, value)


def test_first_visit_categories_are_real_ids(client, db, tenant, client_admin):
    """Regression: the standard categories seeded by the first page view
    were never committed, so the first expense posted was refused with
    'Pick a category'. Use only what the page itself offers."""
    import re

    _on(db, tenant, "expenses")
    _login(client, "admin@acme.example.com", "ClientSecret123")
    page = client.get("/expenses/new").data.decode()
    category_id = re.search(r'<option value="(\d+)"[^>]*>Electricity</option>', page).group(1)
    resp = client.post("/expenses/new", data={"expense_date": TODAY.isoformat(), "category_id": category_id,
                                              "payee_name": "EB", "amount": "900", "payment_mode": "cash"})
    assert resp.status_code == 302, resp.data
    assert Expense.query.one().category.name == "Electricity"
